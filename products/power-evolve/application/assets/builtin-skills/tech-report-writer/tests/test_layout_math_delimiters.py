"""Pagination copies retain native delimiter meanings for sequential OMML importers."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.oxml import parse_xml
from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import catalog_cache

M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
NAMESPACES = 'xmlns:m="' + M[1:-1] + '" xmlns:w="' + W[1:-1] + '"'


def delimiter(begin, end, contents):
    return ('<m:d><m:dPr><m:begChr m:val="' + begin + '"/><m:endChr m:val="' + end
            + '"/><m:sepChr m:val=""/><m:grow/><m:shp m:val="match"/>'
            + '<m:ctrlPr><w:rPr><w:sz w:val="24"/></w:rPr></m:ctrlPr></m:dPr>'
            + '<m:e>' + contents + '</m:e></m:d>')


def omml(contents):
    return parse_xml('<m:oMath ' + NAMESPACES + '>' + contents + '</m:oMath>')


def signatures(root):
    return [{child.tag: etree.tostring(child) for child in properties}
            for properties in root.iter(M + 'dPr')]


class LayoutMathDelimiters(unittest.TestCase):
    def test_bars_and_open_cases_keep_values_and_native_structure(self):
        nested = delimiter('|', '|', '<m:r><m:t>x−y</m:t></m:r>')
        matrix = '<m:m><m:mr><m:e><m:r><m:t>1</m:t></m:r></m:e></m:mr></m:m>'
        root = omml(delimiter('{', '', matrix) + delimiter('(', ')', nested))
        before_properties = signatures(root)
        before_arguments = [etree.tostring(node) for node in root.iter(M + 'e')
                            if next(node.iter(M + 'dPr'), None) is None]
        catalog_cache._order_layout_math_delimiter_properties(root)
        self.assertEqual(signatures(root), before_properties, 'all property values and formatting survive')
        for properties in root.iter(M + 'dPr'):
            self.assertEqual([child.tag for child in properties][:3],
                             [M + 'begChr', M + 'sepChr', M + 'endChr'])
        self.assertEqual([etree.tostring(node) for node in root.iter(M + 'e')
                          if next(node.iter(M + 'dPr'), None) is None], before_arguments)
        self.assertEqual(len(list(root.iter(M + 'd'))), 3)
        self.assertEqual(len(list(root.iter(M + 'm'))), 1)
        once = etree.tostring(root)
        catalog_cache._order_layout_math_delimiter_properties(root)
        self.assertEqual(etree.tostring(root), once, 'layout normalization is idempotent')

    def test_omitted_defaults_and_other_properties_are_not_invented_or_removed(self):
        root = omml('<m:d><m:dPr><m:grow/></m:dPr><m:e><m:r><m:t>x+y</m:t></m:r>'
                    '</m:e></m:d><m:d><m:dPr><m:begChr m:val="["/><m:endChr m:val="]"/>'
                    '</m:dPr><m:e><m:r><m:t>x</m:t></m:r></m:e></m:d>')
        before = etree.tostring(root)
        catalog_cache._order_layout_math_delimiter_properties(root)
        self.assertEqual(etree.tostring(root), before, 'absent properties retain their original default meanings')

    def test_pagination_passes_only_an_ordered_copy_to_uno(self):
        document = Document()
        for instruction in (' TOC \\o "1-3" ', ' TOC \\f F ', ' TOC \\f T '):
            paragraph = document.add_paragraph()
            paragraph._p.append(parse_xml('<w:r ' + NAMESPACES + '><w:fldChar w:fldCharType="begin"/></w:r>'))
            paragraph._p.append(parse_xml('<w:r ' + NAMESPACES + '><w:instrText xml:space="preserve">'
                                          + instruction + '</w:instrText></w:r>'))
            paragraph._p.append(parse_xml('<w:r ' + NAMESPACES + '><w:fldChar w:fldCharType="separate"/></w:r>'))
            paragraph._p.append(parse_xml('<w:r ' + NAMESPACES + '><w:fldChar w:fldCharType="end"/></w:r>'))
        document.add_heading('Synthetic native formulas', 1)
        document.add_paragraph('Synthetic formulas retain source bytes.')
        document.add_paragraph()._p.append(omml(delimiter('|', '|', '<m:r><m:t>x−y</m:t></m:r>')))
        document.add_paragraph()._p.append(omml(delimiter('{', '', '<m:m><m:mr><m:e>'
            '<m:r><m:t>1</m:t></m:r></m:e></m:mr></m:m>')))
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'synthetic.docx'
            pdf = Path(directory) / 'synthetic.pdf'
            document.save(source)
            source_bytes = source.read_bytes()
            with ZipFile(source) as archive:
                original_xml = etree.fromstring(archive.read('word/document.xml'))
                original_parts = {name: archive.read(name) for name in archive.namelist()}
            observed = []

            def inspect_copy(command, **_options):
                request = json.loads(Path(command[-1]).read_text())
                self.assertEqual(request['pdf'], str(pdf))
                self.assertNotEqual(Path(request['docx']), source)
                with ZipFile(request['docx']) as archive:
                    layout = etree.fromstring(archive.read('word/document.xml'))
                    self.assertEqual(set(archive.namelist()), set(original_parts))
                    for name, data in original_parts.items():
                        if name != 'word/document.xml':
                            self.assertEqual(archive.read(name), data)
                self.assertEqual(signatures(layout), signatures(original_xml))
                for properties in layout.iter(M + 'dPr'):
                    self.assertEqual([child.tag for child in properties][:3],
                                     [M + 'begChr', M + 'sepChr', M + 'endChr'])
                observed.append(layout)
                return subprocess.CompletedProcess(command, 0, json.dumps(
                    {'success': True, 'pages': {}, 'page_count': 1}), '')

            with patch.object(catalog_cache.subprocess, 'run', side_effect=inspect_copy):
                self.assertEqual(catalog_cache.paginate(source, pdf)['page_count'], 1)
            self.assertEqual(len(observed), 1)
            self.assertEqual(source.read_bytes(), source_bytes, 'pagination must not rewrite the original DOCX')
            self.assertEqual([child.tag for child in next(original_xml.iter(M + 'dPr'))][:3],
                             [M + 'begChr', M + 'endChr', M + 'sepChr'])


if __name__ == '__main__':
    unittest.main(verbosity=2)
