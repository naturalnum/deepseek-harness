"""Actual R7 equations must retain content and fit their fixed Word table grid."""

from copy import deepcopy
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Mm

sys.dont_write_bytecode = True
SKILL = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL / "scripts"))
import format_existing_report as formatter
from markdown_math import normalize_markdown_display_math

FIXTURE = json.loads((SKILL / "tests/fixtures/native-display-formulas.json").read_text(encoding="utf-8"))
FORMULAS = FIXTURE["formulas"]


def formatted_document():
    """Run production marker/layout code with OMML read from the actual failed DOCX."""
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = section.right_margin = Mm(25)
    for formula in FORMULAS:
        for line in normalize_markdown_display_math(formula["markdown"]).splitlines():
            if line.strip():
                document.add_paragraph(line)
        document.add_paragraph("式(" + formula["number"] + ")后正文保持。")
    native = [parse_xml(formula["omml"]) for formula in FORMULAS]
    with patch.object(formatter, "_pandoc_math_elements", return_value=deepcopy(native)) as converter:
        formatter._replace_formula_markers(document)
    if converter.call_args.args[0] != [formula["latex"] for formula in FORMULAS]:
        raise AssertionError("Reviewed formula meaning changed before layout")
    return document


def legacy_document():
    """Restore the actual released grid mismatch on native, marker-free equations."""
    document = formatted_document()
    for formula, table in zip(FORMULAS, document.tables):
        for column, width in zip(table._tbl.tblGrid, formula["originalGridTwips"]):
            column.set(qn("w:w"), str(width))
        table._tbl.tblPr.find(qn("w:tblW")).set(qn("w:type"), "auto")
        table._tbl.tblPr.find(qn("w:tblW")).set(qn("w:w"), "0")
    return document


class DisplayFormulaLayout(unittest.TestCase):
    def readback(self):
        with tempfile.TemporaryDirectory(prefix="word-formula-") as directory:
            path = Path(directory) / "five-equations.docx"
            formatted_document().save(path)
            return Document(path)

    def test_actual_five_equations_have_a_consistent_usable_fixed_grid(self):
        document = self.readback()
        body_width = document.sections[0].page_width - document.sections[0].left_margin - document.sections[0].right_margin
        self.assertEqual(len(document.tables), 5)
        for formula, table in zip(FORMULAS, document.tables):
            with self.subTest(number=formula["number"]):
                grid = [int(column.get(qn("w:w"))) for column in table._tbl.tblGrid]
                cell_widths = [int(cell._tc.tcPr.find(qn("w:tcW")).get(qn("w:w"))) for cell in table.rows[0].cells]
                self.assertEqual(grid, cell_widths, "Word readers must see the same grid and cell content widths")
                self.assertGreater(table.columns[1].width, 0.8 * body_width,
                                   "The actual cases/median/precision-recall equations must use the body width, not one third")
                total_width = table._tbl.tblPr.find(qn("w:tblW"))
                self.assertEqual(total_width.get(qn("w:type")), "dxa")
                self.assertEqual(int(total_width.get(qn("w:w"))), sum(grid))
                self.assertLessEqual(table.columns[2].width, 0.1 * body_width)

    def test_five_equation_contents_numbers_and_following_prose_survive_roundtrip(self):
        document = self.readback()
        texts = []
        for formula, table in zip(FORMULAS, document.tables):
            math = list(table._tbl.iter(qn("m:oMath")))
            self.assertEqual(len(math), 1)
            text = "".join(node.text or "" for node in math[0].iter(qn("m:t")))
            self.assertEqual(text, formula["mathText"])
            self.assertEqual(table.rows[0].cells[2].text, "（" + formula["number"] + "）")
            texts.append(text)
        self.assertIn("为奇数", texts[0])
        self.assertIn("为偶数", texts[0])
        self.assertIn("i=1,2,…,n", texts[1])
        self.assertIn("recall=TPTP+FN", texts[4])
        self.assertEqual([paragraph.text for paragraph in document.paragraphs],
                         ["式(" + formula["number"] + ")后正文保持。" for formula in FORMULAS])

    def test_existing_marker_free_docx_layout_keeps_data_tables(self):
        legacy = legacy_document()
        data_table = legacy.add_table(rows=1, cols=3)
        data_table.rows[0].cells[1].add_paragraph("保留普通指标表：")
        data_table.rows[0].cells[1].paragraphs[0]._p.append(parse_xml(FORMULAS[0]["omml"]))
        data_table.rows[0].cells[2].text = "（2-1）"
        original_data_grid = [column.get(qn("w:w")) for column in data_table._tbl.tblGrid]
        with tempfile.TemporaryDirectory(prefix="word-formula-existing-") as directory:
            source, destination = Path(directory) / "legacy.docx", Path(directory) / "repaired.docx"
            legacy.save(source)
            before = source.read_bytes()
            reopened = Document(source)
            self.assertEqual(formatter._repair_existing_formula_tables(reopened), 5)
            reopened.save(destination)
            self.assertEqual(source.read_bytes(), before, "format repair must save a new document")
            document = Document(destination)
        for formula, table in zip(FORMULAS, document.tables[:5]):
            with self.subTest(number=formula["number"]):
                self.assertGreater(table.columns[1].width.mm, 140)
                self.assertEqual("".join(node.text or "" for node in table._tbl.iter(qn("m:t"))), formula["mathText"])
                self.assertEqual(table.rows[0].cells[2].text, "（" + formula["number"] + "）")
        self.assertEqual([column.get(qn("w:w")) for column in document.tables[-1]._tbl.tblGrid], original_data_grid)
        self.assertIn("保留普通指标表：", document.tables[-1].rows[0].cells[1].text)


if __name__ == "__main__":
    unittest.main(verbosity=2)
