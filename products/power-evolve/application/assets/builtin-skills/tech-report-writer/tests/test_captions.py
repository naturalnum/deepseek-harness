"""Synthetic DOCX caption marking/formatting/validation; no UNO or user report runs."""
from io import BytesIO
from pathlib import Path
import re
import sys
import tempfile
import unittest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from format_existing_report import _configure_body, _configure_toc, _normalize_inline_symbol_formatting
from report_validator import validate_report
from toc_styles import mark_figure_table_captions

BODY_REFERENCES = (
    "表2-1 中的计数与指标均来自本机对合成样本的实际执行结果，",
    "表2-2 中，符号边界样本集含三个样本，其有符号偏差 a−e",
)


def add_field(paragraph, instruction):
    for kind in ("begin", "instruction", "separate", "end"):
        run = OxmlElement("w:r")
        node = OxmlElement("w:instrText" if kind == "instruction" else "w:fldChar")
        if kind == "instruction":
            node.text = instruction
        else:
            node.set(qn("w:fldCharType"), kind)
        run.append(node)
        paragraph._p.append(run)


def has_tc(paragraph, identifier):
    instruction = "".join(node.text or "" for node in paragraph._p.iter(qn("w:instrText")))
    return re.search(r"\bTC\b.*\\f\s+" + identifier + r"(?:\s|$)", instruction) is not None


def add_figure(document):
    stream = BytesIO()
    Image.new("RGB", (4, 4), "white").save(stream, "PNG")
    stream.seek(0)
    document.add_picture(stream)


class Captions(unittest.TestCase):
    def caption_errors(self, document):
        # The minimal DOCX exercises the public validator's caption checks;
        # unrelated cover/section/page requirements are outside this fixture.
        with tempfile.TemporaryDirectory(prefix="report-captions-") as directory:
            path = Path(directory) / "synthetic.docx"
            document.save(path)
            result = validate_report(str(path))
            readback = Document(path)
        return [error for error in result["errors"]
                if error.startswith(("图题“", "表题“", "非图题段落", "非表题段落"))], readback

    def test_body_references_keep_text_and_body_format_across_catalog_and_body_passes(self):
        document = Document()
        add_field(document.add_paragraph(), ' TOC \\o "1-3" \\h \\z \\u ')
        document.add_paragraph("1 合成验证", style="Heading 1")
        table_caption = document.add_paragraph("表2-1 计数与指标")
        document.add_table(rows=1, cols=2)
        prose = []
        for text in BODY_REFERENCES:
            paragraph = document.add_paragraph()
            paragraph.add_run(text[:5])
            paragraph.add_run(text[5:])
            prose.append(paragraph)
        document.add_paragraph("后续正文保留普通段落格式。")
        add_figure(document)
        figure_caption = document.add_paragraph("图2-1 合成框图")
        document.add_paragraph("图2-1 中的连线用于本段正文说明。")
        document.add_paragraph("不含图表的后续正文。")
        _configure_toc(document)
        _configure_body(document)
        _normalize_inline_symbol_formatting(document)
        for paragraph, text in zip(prose, BODY_REFERENCES):
            self.assertEqual(paragraph.text, text)
            self.assertEqual(paragraph.style.name, "Normal")
            self.assertEqual(paragraph.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
            self.assertEqual(paragraph.runs[0].font.size.pt, 12)
            self.assertFalse(has_tc(paragraph, "T"))
        self.assertTrue(has_tc(table_caption, "T"))
        self.assertTrue(has_tc(figure_caption, "F"))
        self.assertEqual(mark_figure_table_captions(document), {"figures": 1, "tables": 1})
        errors, readback = self.caption_errors(document)
        self.assertEqual(errors, [])
        for text in BODY_REFERENCES:
            self.assertIn(text, [paragraph.text for paragraph in readback.paragraphs])

    def test_true_captions_without_tc_are_rejected_in_each_existing_context(self):
        for context in ("style", "center", "table", "figure", "long-table"):
            with self.subTest(context=context):
                document = Document()
                if context == "figure":
                    add_figure(document)
                    document.add_paragraph("")
                title = ("图2-1 合成框图" if context == "figure" else "表2-1 固定输入的验证结果")
                if context == "long-table":
                    title += "及不同严格阈值下的计数和指标" * 20
                paragraph = document.add_paragraph(title, style="Caption" if context == "style" else "Normal")
                if context == "center":
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if context in ("table", "long-table"):
                    document.add_paragraph("")
                    document.add_table(rows=1, cols=1)
                errors, _ = self.caption_errors(document)
                self.assertEqual(len(errors), 1, errors)
                self.assertIn("缺少" + ("图" if context == "figure" else "表") + "目录 TC 域", errors[0])
                self.assertEqual(mark_figure_table_captions(document),
                                 {"figures": int(context == "figure"), "tables": int(context != "figure")})
                self.assertEqual(self.caption_errors(document)[0], [])

    def test_tc_on_ordinary_body_prose_is_rejected_and_catalog_pass_removes_it(self):
        for text, identifier in [(text, "T") for text in BODY_REFERENCES] + [
            ("图2-1 中的连线用于本段正文说明。", "F"),
        ]:
            with self.subTest(text=text):
                document = Document()
                paragraph = document.add_paragraph(text)
                document.add_paragraph("后续正文阻断相邻对象判定。")
                add_field(paragraph, ' TC "' + text + '" \\f ' + identifier + ' \\l "1" ')
                errors, _ = self.caption_errors(document)
                self.assertEqual(len(errors), 1, errors)
                self.assertIn("非" + ("图" if identifier == "F" else "表") + "题段落不得写入", errors[0])
                self.assertEqual(mark_figure_table_captions(document), {"figures": 0, "tables": 0})
                self.assertFalse(has_tc(paragraph, identifier))
                self.assertEqual(paragraph.text, text)
                self.assertEqual(self.caption_errors(document)[0], [])

    def test_continued_table_stays_formatted_without_a_second_catalog_item(self):
        document = Document()
        document.add_paragraph("1 合成验证", style="Heading 1")
        caption = document.add_paragraph("表2-1 固定结果")
        document.add_table(rows=1, cols=1)
        continued = document.add_paragraph("表2-1 （续）")
        document.add_table(rows=1, cols=1)
        self.assertEqual(mark_figure_table_captions(document), {"figures": 0, "tables": 1})
        _configure_body(document)
        self.assertEqual(continued.text, "表2-1 （续）")
        self.assertEqual(continued.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(continued.runs[0].font.size.pt, 10.5)
        self.assertTrue(continued.paragraph_format.keep_with_next)
        self.assertTrue(has_tc(caption, "T"))
        self.assertFalse(has_tc(continued, "T"))
        self.assertEqual(self.caption_errors(document)[0], [])
        add_field(continued, ' TC "表2-1 （续）" \\f T \\l "1" ')
        self.assertIn("非表题段落不得写入", self.caption_errors(document)[0][0])

    def test_existing_explanatory_lead_rejection_is_preserved(self):
        document = Document()
        paragraph = document.add_paragraph("表2-1 给出合成结果。", style="Caption")
        document.add_table(rows=1, cols=1)
        self.assertEqual(mark_figure_table_captions(document), {"figures": 0, "tables": 0})
        with tempfile.TemporaryDirectory(prefix="report-captions-") as directory:
            path = Path(directory) / "explanation.docx"
            document.save(path)
            errors = validate_report(str(path))["errors"]
        self.assertTrue(any("图表说明段不得以图号/表号起句" in error for error in errors))
        self.assertFalse(has_tc(paragraph, "T"))


if __name__ == "__main__":
    unittest.main(verbosity=2)
