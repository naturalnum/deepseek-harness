"""Final reviewed report text must survive deterministic Word preparation.

Run with the report Python dependencies: python3 tests/test_orchestration_adapter.py
"""
import contextlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))
from orchestration_adapter import chapter_data, final_report_data, abstract_data, create_word_report, main as adapter_main
from docx import Document
from engine.parser import parse_html_to_elements
from format_existing_report import _configure_body
from docx.oxml.ns import qn
from docx.shared import Pt
from report_validator import validate_report
from main import _validate_references_data
from report_generator import _build_appendices_html, _build_references_html


FINAL_MARKER = "REPORT-REVISION-20261003"
OLD_MARKER = "REPORT-OLD-20261003"
FINAL_REPORT = f"""# 测试负荷异常识别技术报告

## 中文摘要

中文摘要包含最终确认的测试结论。

**关键词：** 电网；负荷；中位数

## ABSTRACT

The reviewed report records the final test result.

**Keywords:** power grid; load; median

# 1 项目概述

{FINAL_MARKER}：测试数据仅支持局部结论。

## 1.1 研究范围

不得把本测试解释为生产性能。

# 2 配电负荷异常识别

表2-1 测试结果

| 指标 | 数值 |
| --- | ---: |
| 识别数 | 3 |

# 3 总结

确认的总结只包含本测试的结论。

# 参考文献

[1] 测试文献作者. 测试方法说明[R]. 2026.
"""


def payload(source=FINAL_REPORT):
    return {
        "source_report": {"artifact_id": "report-qa", "version": 2,
                          "content": {"markdown": source}},
        "approved_chapters": [{"artifact_id": "chapter-1", "version": 1,
                               "content": {"title": "项目概述", "markdown": f"# 1 项目概述\n\n{OLD_MARKER}"}}],
        "abstract": {"content": {"abstract_cn": OLD_MARKER, "keywords_cn": "旧；词",
                                 "abstract_en": OLD_MARKER, "keywords_en": "old; words"}},
        "reference_plan": {"content": {"items": [{"gbt7714_entry": f"测试作者. {OLD_MARKER}[R]. 2026."}]}},
        "word_metadata": {"report_no": "QA-001", "class_no": "QA-TM", "secret_level": "公开测试",
                          "project_source": "非生产测试项目", "title_cn": "测试负荷异常识别技术报告",
                          "title_en": "Test Load Anomaly Report", "project_name": "测试项目",
                          "report_date": "2026年10月"},
    }


class FinalReportSource(unittest.TestCase):
    def test_explicit_front_contents_does_not_enter_abstracts_or_body(self):
        baseline = final_report_data(payload()["source_report"])
        contents = "\n\nTOC-ONLY-MARKER\n\n- 1 项目概述\n  - 1.1 研究范围\n- 2 配电负荷异常识别\n- 3 总结\n\n"
        for heading in ("# 目录", "## 目录"):
            for position in ("## 中文摘要", "## ABSTRACT", "# 1 项目概述"):
                with self.subTest(heading=heading, position=position):
                    source = FINAL_REPORT.replace(position, heading + contents + position)
                    result = final_report_data(payload(source)["source_report"])
                    self.assertEqual(result, baseline)
                    self.assertNotIn("TOC-ONLY-MARKER", json.dumps(result, ensure_ascii=False))

    def test_contents_must_be_unique_nonempty_and_before_body(self):
        for source in (
            FINAL_REPORT.replace("# 1 项目概述", "# 目录\n\n# 1 项目概述"),
            FINAL_REPORT.replace("# 1 项目概述", "# 目录\n\n- 1 项目概述\n\n## 目录\n\n- 2 研究\n\n# 1 项目概述"),
            *(FINAL_REPORT.replace("# 2 配电", heading + "\n\n- 2 研究\n\n# 2 配电")
              for heading in ("# 目录", "## 目录")),
            FINAL_REPORT + "\n## 目录\n\n- 1 项目概述\n",
        ):
            with self.subTest(source=source):
                with self.assertRaises(ValueError):
                    chapter_data(payload(source))

    def test_bilingual_titles_require_matching_current_word_metadata(self):
        data = payload()
        source = FINAL_REPORT.replace("## 中文摘要", "# " + data["word_metadata"]["title_en"] + "\n\n## 中文摘要")
        baseline = final_report_data(data["source_report"])
        data["source_report"]["content"]["markdown"] = source
        self.assertEqual(final_report_data(data["source_report"], data["word_metadata"]), baseline)
        self.assertEqual(chapter_data(data), baseline["chapters"])
        for metadata in (None, {}, {**data["word_metadata"], "title_cn": "其他中文题名"},
                         {**data["word_metadata"], "title_en": "Other English Title"}):
            with self.subTest(metadata=metadata):
                with self.assertRaises(ValueError):
                    final_report_data(data["source_report"], metadata)

    def test_unknown_extra_titles_and_late_english_title_remain_rejected(self):
        metadata = payload()["word_metadata"]
        for source in (
            FINAL_REPORT.replace("## 中文摘要", "# 未支持的前置节\n\n## 中文摘要"),
            FINAL_REPORT.replace("## 中文摘要", "# " + metadata["title_en"] + "\n\n# 第三个题名\n\n## 中文摘要"),
            FINAL_REPORT.replace("## ABSTRACT", "# " + metadata["title_en"] + "\n\n## ABSTRACT"),
            FINAL_REPORT.replace("# 参考文献", "# 未支持的后置节\n\n# 参考文献"),
        ):
            with self.subTest(source=source):
                with self.assertRaises(ValueError):
                    final_report_data(payload(source)["source_report"], metadata)

    def test_contents_heading_in_fenced_program_remains_program_text(self):
        source = FINAL_REPORT.replace(FINAL_MARKER, "```python\n# 目录\nprint('not a contents section')\n```\n\n" + FINAL_MARKER)
        result = final_report_data(payload(source)["source_report"])
        self.assertIn("# 目录", result["chapters"][0]["content_html"])

    def test_source_explanation_is_preserved_as_appendix_not_reference_entry(self):
        note = "本报告引用本地未出版 QA 材料，不涉及外部 URL 或出版来源。"
        source = FINAL_REPORT + "\n# 附录 A 复现程序\n\n```python\n# 目录\nprint('preserved')\n```\n\n# 附录 B 参考文献来源与获取链接\n\n" + note + "\n"
        result = final_report_data(payload(source)["source_report"])
        self.assertEqual([row["id"] for row in result["appendices"]], ["A", "B"])
        self.assertEqual([row["title"] for row in result["chapters"]], ["项目概述", "配电负荷异常识别", "总结"])
        self.assertIn(note, result["appendices"][1]["content_html"])
        self.assertNotIn(note, result["references"]["references_html"])
        self.assertTrue(_validate_references_data(result["references"])[0])
        document = Document()
        back_matter = _build_appendices_html({"appendices": result["appendices"]}) + _build_references_html(result["references"])
        _, errors, _ = parse_html_to_elements(document, back_matter, "CONTENT")
        self.assertEqual(errors, [])
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source-appendix.docx"
            document.save(path)
            readback = Document(path)
            text = [paragraph.text for paragraph in readback.paragraphs]
            self.assertEqual(text.count(note), 1)
            self.assertIn("附录 B 参考文献来源与获取链接", text)
            self.assertTrue(text.index(note) < text.index("参考文献"))
            errors = validate_report(str(path))["errors"]
            self.assertFalse(any("参考文献" in error for error in errors), errors)
        with self.assertRaisesRegex(ValueError, "unsupported top-level heading"):
            final_report_data(payload(FINAL_REPORT + "\n# 参考文献来源与获取链接\n\n" + note)["source_report"])

    def test_final_heading_numbers_names_and_levels_keep_reviewed_order(self):
        source = FINAL_REPORT.replace(
            "不得把本测试解释为生产性能。",
            "不得把本测试解释为生产性能。\n\n### 1.1.1 测试边界\n\n只描述测试边界。\n\n"
            "## 测试备注\n\n保留未编号标题。\n\n### 测试说明\n\n不补造编号。",
        ).replace("表2-1 测试结果", "## 2.1 固定测试结果\n\n### 2.1.1 计算方法\n\n表2-1 测试结果")
        final = final_report_data(payload(source)["source_report"])
        document = Document()
        for chapter in final["chapters"]:
            _, errors, _ = parse_html_to_elements(document, chapter["content_html"], "CONTENT")
            self.assertEqual(errors, [])
        _configure_body(document)
        expected = [("Heading 1", "1 项目概述"), ("Heading 2", "1.1 研究范围"),
                    ("Heading 3", "1.1.1 测试边界"), ("Heading 2", "测试备注"),
                    ("Heading 3", "测试说明"), ("Heading 1", "2 配电负荷异常识别"),
                    ("Heading 2", "2.1 固定测试结果"), ("Heading 3", "2.1.1 计算方法"),
                    ("Heading 1", "3 总结")]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "reviewed-headings.docx"
            document.save(path)
            readback = Document(path)
            headings = [paragraph for paragraph in readback.paragraphs
                        if paragraph.style.name.startswith("Heading ")]
            self.assertEqual([(paragraph.style.name, paragraph.text) for paragraph in headings], expected)
            for paragraph in headings:
                self.assertFalse(paragraph._p.xpath("./w:pPr/w:numPr"), "the source already supplies numbers")
                self.assertEqual(paragraph._p.pPr.find(qn("w:outlineLvl")).get(qn("w:val")),
                                 str(int(paragraph.style.name[-1]) - 1))
            errors = validate_report(str(path))["errors"]
            self.assertFalse(any("级标题“" in error for error in errors), errors)
            self.assertFalse(any("正文段内符号 run“" in error for error in errors), errors)
            unnumbered = next(paragraph for paragraph in headings if paragraph.text == "测试备注")
            for run in unnumbered.runs:
                run.font.size = Pt(12)
            readback.save(path)
            errors = validate_report(str(path))["errors"]
            self.assertIn("二级标题“测试备注”格式不符合表 B-1", errors)

    def test_numeric_prose_is_not_a_list_or_heading(self):
        prose = "1.1 这里是正文数据说明。"
        source = FINAL_REPORT.replace("不得把本测试解释为生产性能。", prose + "\n\n1. 真实列表项")
        final = final_report_data(payload(source)["source_report"])
        document = Document()
        for chapter in final["chapters"]:
            _, errors, _ = parse_html_to_elements(document, chapter["content_html"], "CONTENT")
            self.assertEqual(errors, [])
        _configure_body(document)
        paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text == prose]
        self.assertEqual(len(paragraphs), 1, "keep the complete numeric prose")
        paragraph = paragraphs[0]
        self.assertEqual(paragraph.style.name, "Normal")
        self.assertFalse(paragraph._p.xpath("./w:pPr/w:numPr"))
        self.assertEqual(paragraph.runs[0].font.size.pt, 12)
        self.assertFalse(paragraph.runs[0].font.bold)
        self.assertEqual(paragraph._p.pPr.find(qn("w:snapToGrid")).get(qn("w:val")), "0")
        list_item = next(paragraph for paragraph in document.paragraphs if paragraph.text == "真实列表项")
        self.assertEqual(list_item.style.name, "List Number")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "numeric-prose.docx"
            document.save(path)
            errors = validate_report(str(path))["errors"]
            self.assertFalse(any(f"二级标题“{prose}" in error for error in errors), errors)

    def test_final_report_replaces_prior_confirmed_chapter_text(self):
        chapters = chapter_data(payload())
        text = "\n".join(row["content_html"] for row in chapters)
        self.assertIn(FINAL_MARKER, text)
        self.assertNotIn(OLD_MARKER, text)

    def test_front_matter_and_references_are_not_body_chapters(self):
        chapters = chapter_data(payload())
        self.assertEqual([row["title"] for row in chapters],
                         ["项目概述", "配电负荷异常识别", "总结"])
        self.assertNotIn("ABSTRACT", "\n".join(row["content_html"] for row in chapters))
        self.assertNotIn("测试文献作者", "\n".join(row["content_html"] for row in chapters))

    def test_absent_final_source_keeps_standalone_chapter_support(self):
        data = payload()
        del data["source_report"]
        self.assertIn(OLD_MARKER, chapter_data(data)[0]["content_html"])

    def test_reviewed_abstracts_references_and_table_reach_word_elements(self):
        data = final_report_data(payload()["source_report"])
        abstracts = abstract_data(data["abstract"])
        self.assertIn("最终确认", abstracts["abstract_cn"])
        self.assertIn("reviewed report", abstracts["abstract_en"])
        self.assertEqual(abstracts["keywords_cn"], "电网；负荷；中位数")
        self.assertIn("测试文献作者", data["references"]["references_html"])
        document = Document()
        fragments = [abstracts["abstract_cn"], abstracts["abstract_en"],
                     *(chapter["content_html"] for chapter in data["chapters"]),
                     data["references"]["references_html"]]
        for fragment in fragments:
            _, errors, _ = parse_html_to_elements(document, fragment, "CONTENT")
            self.assertEqual(errors, [])
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "final.docx"
            document.save(path)
            readback = Document(path)
            text = "\n".join(paragraph.text for paragraph in readback.paragraphs)
            self.assertIn(FINAL_MARKER, text)
            self.assertNotIn(OLD_MARKER, text)
            self.assertIn("reviewed report", text)
            self.assertIn("测试文献作者", text)
            self.assertEqual(len(readback.tables), 1)
            self.assertEqual([[cell.text for cell in row.cells] for row in readback.tables[0].rows],
                             [["指标", "数值"], ["识别数", "3"]])

    def test_file_source_and_supported_front_matter_heading_forms(self):
        source = FINAL_REPORT.replace("## 中文摘要", "# 摘要").replace("## ABSTRACT", "# Abstract")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.md"
            path.write_text("\ufeff" + source, encoding="utf-8")
            data = payload()
            data["source_report"]["content"] = {"markdown_path": str(path)}
            self.assertEqual(len(chapter_data(data)), 3)

    def test_create_receipt_uses_only_this_edition_data_directory(self):
        # Exercise the adapter's JSON and receipt paths with the real Word element
        # writer; Linux UNO template rendering is covered by container acceptance.
        def command(cmd, env, cwd):
            if cmd[-1] != "create-report":
                return subprocess.run(cmd, env={**env, "PYTHONDONTWRITEBYTECODE": "1"}, cwd=cwd,
                                      capture_output=True, text=True, check=True, timeout=20)
            directory = Path(env["TECH_REPORT_OUTPUT_DIR"])
            document = Document()
            abstracts = json.loads((directory / "abstract.json").read_text())
            for key in ("abstract_cn", "abstract_en"):
                _, errors, _ = parse_html_to_elements(document, abstracts[key], "CONTENT")
                self.assertEqual(errors, [])
            for path in sorted(directory.glob("chapter_*.json")):
                row = json.loads(path.read_text())
                _, errors, _ = parse_html_to_elements(document, row["content_html"], "CONTENT")
                self.assertEqual(errors, [])
            references = json.loads((directory / "references.json").read_text())
            _, errors, _ = parse_html_to_elements(document, references["references_html"], "CONTENT")
            self.assertEqual(errors, [])
            document.save(directory / "report.docx")
            return subprocess.CompletedProcess(cmd, 0, "", "")

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output = root / "word"
            output.mkdir()
            stale = output / "chapter_9.json"
            stale.write_text(OLD_MARKER)
            stale_appendices = output / "appendices.json"
            stale_appendices.write_text(OLD_MARKER)
            data = payload()
            source = FINAL_REPORT.replace("## 中文摘要", "# " + data["word_metadata"]["title_en"] + "\n\n## 中文摘要")
            source = source.replace("# 1 项目概述", "## 目录\n\nTOC-ONLY-MARKER\n\n- 1 项目概述\n- 2 配电负荷异常识别\n- 3 总结\n\n# 1 项目概述")
            data["source_report"]["content"]["markdown"] = source
            data["output_preferences"] = {"workspace_dir": str(root), "output_dir": str(output)}
            request = {"skill_name": "tech-report-writer", "action": "create_word_report", "payload": data,
                       "orchestration": {"enabled": True, "request_id": "source-qa",
                                         "caller_skill": "report-workflow-orchestrator"}}
            with patch("orchestration_adapter.run_command", side_effect=command):
                response = create_word_report(request)
            content = response["artifact"]["content"]
            current = Path(content["data_dir"])
            self.assertEqual(current.parent, output.resolve())
            self.assertEqual(Path(content["docx_path"]).parent, current)
            self.assertEqual(stale.read_text(), OLD_MARKER)
            self.assertEqual(stale_appendices.read_text(), OLD_MARKER)
            self.assertFalse((current / "chapter_9.json").exists())
            self.assertFalse((current / "appendices.json").exists())
            self.assertEqual(content["source_report_version"], 2)
            document = Document(content["docx_path"])
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn(FINAL_MARKER, text)
            self.assertNotIn(OLD_MARKER, text)
            self.assertIn("最终确认", text)
            self.assertIn("reviewed report", text)
            self.assertIn("测试文献作者", text)
            self.assertNotIn("TOC-ONLY-MARKER", text)
            self.assertNotIn(data["word_metadata"]["title_en"], text)
            self.assertEqual(len(document.tables), 1)

            del data["source_report"]
            with patch("orchestration_adapter.run_command", side_effect=command):
                legacy = create_word_report(request)["artifact"]["content"]
            self.assertNotEqual(legacy["data_dir"], content["data_dir"])
            legacy_text = "\n".join(paragraph.text for paragraph in Document(legacy["docx_path"]).paragraphs)
            self.assertIn(OLD_MARKER, legacy_text)
            self.assertNotIn(FINAL_MARKER, legacy_text)
            self.assertEqual(Path(legacy["docx_path"]).parent, Path(legacy["data_dir"]))

    def test_broken_final_sources_cannot_select_old_chapters(self):
        for source in (None, "", "Unstructured report", FINAL_REPORT.replace("# 2 配电", "# 4 配电"),
                       FINAL_REPORT.replace("# 1 项目概述", "# 1项目概述"),
                       FINAL_REPORT.replace("## ABSTRACT", "## OTHER"),
                       FINAL_REPORT.replace("| 识别数 | 3 |", "| 识别数 | 3 | 4 |")):
            with self.subTest(source_type=type(source).__name__):
                with self.assertRaises(ValueError):
                    chapter_data(payload(source))

    def test_invalid_final_source_is_blocked_without_falling_back(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            data = payload("This is not a final report")
            data["output_preferences"] = {"output_dir": str(root / "output")}
            request = {"skill_name": "tech-report-writer", "action": "create_word_report", "payload": data,
                       "orchestration": {"enabled": True, "request_id": "source-qa",
                                         "caller_skill": "report-workflow-orchestrator"}}
            request_path, response_path = root / "request.json", root / "response.json"
            request_path.write_text(json.dumps(request), encoding="utf-8")
            result = subprocess.run([sys.executable, "-B", str(SCRIPTS / "orchestration_adapter.py"),
                                     str(request_path), str(response_path)], capture_output=True, text=True, timeout=20,
                                    env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"})
            self.assertEqual(result.returncode, 1, result.stdout + result.stderr)
            response = json.loads(response_path.read_text(encoding="utf-8"))
            self.assertEqual(response["status"], "blocked")
            self.assertIn("final source_report", response["error"]["message"])
            self.assertEqual(list(root.rglob("*.docx")), [])


class CommandFailures(unittest.TestCase):
    def request(self, root):
        data = payload()
        data["source_report"]["content"]["markdown"] = FINAL_REPORT.replace(
            "## 中文摘要", "# " + data["word_metadata"]["title_en"] + "\n\n## 中文摘要")
        data["output_preferences"] = {"workspace_dir": str(root), "output_dir": str(root / "word")}
        return {"skill_name": "tech-report-writer", "action": "create_word_report", "payload": data,
                "orchestration": {"enabled": True, "request_id": "diagnostic-qa",
                                  "caller_skill": "report-workflow-orchestrator"}}

    def invoke(self, root, request, command):
        request_path, response_path = root / "request.json", root / "response.json"
        request_path.write_text(json.dumps(request), encoding="utf-8")
        visible = io.StringIO()
        with patch("orchestration_adapter.run_command", side_effect=command), \
                patch("sys.argv", ["orchestration_adapter.py", str(request_path), str(response_path)]), \
                contextlib.redirect_stdout(visible):
            code = adapter_main()
        response = json.loads(response_path.read_text(encoding="utf-8"))
        self.assertEqual(json.loads(visible.getvalue()), response)
        self.assertEqual(code, 1)
        self.assertEqual(response["request_id"], "diagnostic-qa")
        self.assertEqual(response["status"], "blocked")
        self.assertNotIn("artifact", response)
        self.assertNotIn("validation_result", response)
        return response, visible.getvalue()

    @staticmethod
    def run_child(cmd, env, cwd):
        return subprocess.run(cmd, env={**env, "PYTHONDONTWRITEBYTECODE": "1"}, cwd=cwd,
                              capture_output=True, text=True, check=True, timeout=20)

    def test_created_partial_docx_remains_blocked_with_bounded_actual_logs(self):
        stdout = "输出🙂" * 1700 + "\nCREATE-END\n"
        stderr = "TRACE-START\n" + "原因🙂" * 1700 + (
            "\nRuntimeError: 目录刷新失败: UNO pagination failed: "
            "LibreOffice Math is required to render native Word formulas\n")
        # The actual child saves a DOCX and exits nonzero; it substitutes for UNO,
        # while the adapter and preflight CLI retain their real execution paths.
        script = ("import os, sys\nfrom pathlib import Path\nfrom docx import Document\n"
                  "doc = Document(); doc.add_paragraph('PARTIAL-DOCX')\n"
                  "doc.save(Path(os.environ['TECH_REPORT_OUTPUT_DIR']) / 'partial.docx')\n"
                  f"sys.stdout.write({stdout!r}); sys.stderr.write({stderr!r}); sys.exit(1)\n")
        commands = []

        def command(cmd, env, cwd):
            commands.append(cmd[2:])
            return self.run_child([sys.executable, "-c", script] if cmd[2] == "create-report" else cmd, env, cwd)

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            response, visible = self.invoke(root, self.request(root), command)
            error = response["error"]
            self.assertEqual(error["phase"], "create-report")
            self.assertEqual(error["returncode"], 1)
            data_dir = Path(error["data_dir"])
            self.assertEqual(data_dir.parent, (root / "word").resolve())
            self.assertEqual(Document(data_dir / "partial.docx").paragraphs[0].text, "PARTIAL-DOCX")
            self.assertEqual(commands[-1], ["create-report"])
            self.assertTrue(any(cmd == ["validate-chapter", "2"] for cmd in commands))
            for key, actual in (("stdout", stdout), ("stderr", stderr)):
                self.assertLessEqual(len(error[key].encode("utf-8")), 4096)
                self.assertEqual(error[f"{key}_bytes"], len(actual.encode("utf-8")))
                self.assertTrue(error[f"{key}_truncated"])
                self.assertTrue(actual.endswith(error[key]))
            self.assertTrue(error["stdout"].endswith("CREATE-END\n"))
            self.assertIn("LibreOffice Math is required", error["stderr"])
            self.assertNotIn("TRACE-START", error["stderr"])
            self.assertLess(len(visible.encode("utf-8")), 11000)

    def test_preflight_failure_keeps_specific_phase_and_never_creates_docx(self):
        commands = []

        def command(cmd, env, cwd):
            commands.append(cmd[2:])
            failed = cmd[2:] == ["validate-chapter", "2"]
            actual = [sys.executable, "-c", "import sys; sys.stderr.write('chapter rejected\\n'); sys.exit(3)"] if failed else cmd
            return self.run_child(actual, env, cwd)

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            response, _ = self.invoke(root, self.request(root), command)
            error = response["error"]
            self.assertEqual(error["phase"], "validate-chapter 2")
            self.assertEqual(error["returncode"], 3)
            self.assertEqual(error["stdout"], "")
            self.assertEqual(error["stdout_bytes"], 0)
            self.assertEqual(error["stderr"], "chapter rejected\n")
            self.assertFalse(error["stdout_truncated"])
            self.assertFalse(error["stderr_truncated"])
            self.assertTrue((Path(error["data_dir"]) / "chapter_2.json").is_file())
            self.assertEqual(commands[-1], ["validate-chapter", "2"])
            self.assertNotIn(["create-report"], commands)
            self.assertEqual(list(root.rglob("*.docx")), [])

    def test_format_failure_retains_diagnostics_and_original_docx(self):
        def command(cmd, env, cwd):
            self.assertEqual(Path(cmd[1]).name, "format_existing_report.py")
            script = ("import sys\nfrom docx import Document\n"
                      "doc = Document(); doc.add_paragraph('PARTIAL-FORMAT'); doc.save(sys.argv[1])\n"
                      "sys.stdout.write('format started\\n'); sys.stderr.write('format rejected\\n'); sys.exit(5)\n")
            return self.run_child([sys.executable, "-c", script, cmd[-1]], env, cwd)

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            src = root / "original.docx"
            document = Document(); document.add_paragraph("ORIGINAL"); document.save(src)
            original = src.read_bytes()
            request = self.request(root)
            request["action"] = "format_existing_word_report"
            request["payload"] = {"word_report": {"artifact_id": "word-qa", "version": 1,
                                                   "content": {"docx_path": str(src)}}}
            response, _ = self.invoke(root, request, command)
            error = response["error"]
            self.assertEqual(error["phase"], "format-existing-report")
            self.assertEqual(error["returncode"], 5)
            self.assertEqual(error["data_dir"], str(root.resolve()))
            self.assertEqual(error["stdout"], "format started\n")
            self.assertEqual(error["stderr"], "format rejected\n")
            self.assertEqual(src.read_bytes(), original)
            self.assertEqual(Document(root / "original_格式修复版.docx").paragraphs[0].text, "PARTIAL-FORMAT")

    def test_non_subprocess_error_preserves_original_message(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            request = self.request(root)
            del request["payload"]["word_metadata"]["title_cn"]
            command = Mock(side_effect=AssertionError("must not run a child"))
            response, _ = self.invoke(root, request, command)
            self.assertEqual(response["error"], {"message": "missing required Word metadata: title_cn"})
            command.assert_not_called()
            self.assertFalse((root / "word").exists())


if __name__ == "__main__":
    unittest.main(verbosity=2)
