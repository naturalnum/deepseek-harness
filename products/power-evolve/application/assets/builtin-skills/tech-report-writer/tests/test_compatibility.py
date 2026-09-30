"""Container-only report regressions. Requires the qa2 Linux UNO toolchain.

Run: python3 /skills/tech-report-writer/tests/test_compatibility.py
Set REPORT_COMPAT_OUTPUT to retain the generated DOCX/PDF/PNG evidence.
"""
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))
from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw
from catalog_cache import _load, _save, _ranges, W, refresh_catalogs
from format_existing_report import _cover_boundaries, format_report
from render_for_visual_qa import render
from report_validator import validate_report


class ReportCompatibility(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = Path(os.environ.get("REPORT_COMPAT_OUTPUT", tempfile.mkdtemp(prefix="report-compat-")))
        cls.root.mkdir(parents=True, exist_ok=True)
        cls.documents = []
        for index, source in enumerate(("测试样例", "咨询委托单位", "")):
            output = cls.root / str(index)
            output.mkdir(exist_ok=True)
            cover = dict(report_no="QA-001", class_no="TM", secret_level="公开",
                project_source=source, title_cn="技能兼容性报告", title_en="Skill Compatibility Report",
                project_name="最小测试项目", series_no="", cooperator=[], report_date="2026年9月",
                work_period="2026年9月", leader="张三", staff="李四", writer="王五")
            payloads = {
                "cover_info": cover,
                "abstract": dict(abstract_cn="本报告检验技能脚本的运行兼容性。", keywords_cn="兼容性；测试",
                    abstract_en="This report verifies runtime compatibility.", keywords_en="compatibility; test"),
                "outline": dict(chapters=[dict(num=1, title="样例正文", sections=[])]),
                "chapter_1": dict(num=1, title="样例正文", content_html='<h1>1 样例正文</h1><p>中文字体与报告模板兼容性验证。</p>'),
            }
            for name, data in payloads.items():
                (output / (name + ".json")).write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
            process = subprocess.run([sys.executable, str(SCRIPTS / "main.py"), "create-report"],
                env={**os.environ, "TECH_REPORT_OUTPUT_DIR": str(output)}, text=True, capture_output=True)
            if process.returncode:
                raise AssertionError(process.stdout + process.stderr)
            cls.documents.append(next(output.glob("*.docx")))

    def test_arbitrary_and_blank_cover_sources(self):
        for path in self.documents:
            self.assertTrue(validate_report(str(path))["valid"])
            source, _, _, _ = _cover_boundaries(Document(path))
            self.assertEqual(Document(path).paragraphs[source].text.strip(),
                json.loads((path.parent / "cover_info.json").read_text())["project_source"])

    def test_missing_cover_data_fails_clearly(self):
        output = self.root / "missing"
        output.mkdir(exist_ok=True)
        (output / "cover_info.json").write_text('{}')
        process = subprocess.run([sys.executable, str(SCRIPTS / "main.py"), "create-report"],
            env={**os.environ, "TECH_REPORT_OUTPUT_DIR": str(output)}, text=True, capture_output=True)
        self.assertNotEqual(process.returncode, 0)
        self.assertIn("title_cn", process.stdout + process.stderr)
        broken = Document()
        with self.assertRaisesRegex(ValueError, "工作单位"):
            _cover_boundaries(broken)

    def test_main_catalog_renders_and_empty_or_wrong_pages_fail(self):
        good = self.documents[0]
        result = render(good, self.root / "good-render")
        (self.root / "good-render.json").write_text(json.dumps(result, ensure_ascii=False, indent=2))
        self.assertTrue(result["success"], result)
        self.assertEqual(result["catalog_qa"]["main_cached_entries"], 3)
        for case in ("blank", "wrong-page"):
            path = self.root / (case + ".docx")
            files, root = _load(good)
            _, nodes = _ranges(root)["C"]
            if case == "blank":
                for node in nodes:
                    for text in node.iter(W + "t"):
                        text.text = ""
            else:
                list(nodes[-1].iter(W + "t"))[-1].text = "999"
            _save(path, files, root)
            failed = render(path, self.root / (case + "-render"))
            self.assertFalse(failed["success"], failed)
            self.assertFalse(failed["catalog_qa"]["ok"], failed)
            self.assertTrue(failed["catalog_qa"]["incorrect_catalog_pages"], failed)
            (self.root / (case + "-render.json")).write_text(json.dumps(failed, ensure_ascii=False, indent=2))

    def test_multi_page_toc_and_custom_tc_survive(self):
        path = self.root / "long.docx"
        document = Document(self.documents[0])
        picture = self.root / "diagram.png"
        image = Image.new("RGB", (600, 120), "white")
        ImageDraw.Draw(image).rectangle((5, 5, 590, 110), outline="black", width=2)
        image.save(picture)
        document.add_picture(str(picture), width=Inches(3))
        document.add_paragraph("图1-1 兼容性框架")
        document.add_paragraph("表1-1 兼容性记录")
        table = document.add_table(rows=2, cols=2)
        for cell, value in zip([cell for row in table.rows for cell in row.cells], ["项目", "结果", "目录", "通过"]):
            cell.text = value
        for index in range(1, 61):
            document.add_paragraph(f"1.{index} 回归章节{index}", style="Heading 2")
            document.add_paragraph("用于检验目录跨页后的实际页码，不含业务结果。")
        document.save(path)
        formatted = self.root / "long-formatted.docx"
        format_report(path, formatted)
        structural = validate_report(str(formatted))
        self.assertTrue(structural["valid"], structural)
        _, root = _load(formatted)
        instructions = " ".join(n.text or "" for n in root.iter(W + "instrText"))
        for marker in ('TC "图1-1 兼容性框架" \\f F', 'TC "表1-1 兼容性记录" \\f T'):
            self.assertIn(marker, instructions)
        result = render(formatted, self.root / "long-render")
        (self.root / "long-render.json").write_text(json.dumps(result, ensure_ascii=False, indent=2))
        self.assertTrue(result["success"], result)
        self.assertGreater(result["catalog_qa"]["main_cached_entries"], 60)
        start, end = result["catalog_qa"]["catalog_physical_pages"]
        self.assertGreater(end, start)
        before = formatted.read_bytes()
        refresh_catalogs(formatted)
        self.assertEqual(formatted.read_bytes(), before, "refresh must converge without accumulating entries")


if __name__ == "__main__":
    unittest.main(verbosity=2)
