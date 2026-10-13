"""The report's declared line spacing must not snap to a legacy page grid."""
from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from config import REPORT_WRITING_UNIT
from format_existing_report import _configure_front_matter, _format_body, _format_reference


def abstract_fixture():
    document = Document()
    for text in ['报告编号：SYNTHETIC', '分类号：TM', '密级：公开', '合成测试委托',
                 '合成摘要网格检查', 'Synthetic Abstract Grid Check', '《合成项目》',
                 REPORT_WRITING_UNIT, '2026年10月']:
        document.add_paragraph(text)
    document.add_section(WD_SECTION_START.NEW_PAGE)
    for text in ['工作单位：' + REPORT_WRITING_UNIT, '工作时间：2026年10月',
                 '项目负责人：测试人员', '工作人员：合成人员', '报告编写：合成作者']:
        document.add_paragraph(text)
    document.add_section(WD_SECTION_START.NEW_PAGE)
    expected = {}
    for title, texts, keywords, east_asia in (
        ('内容摘要', ['本合成摘要只检验 Word 段落布局，保留编号 A1 与单位 kW。',
                    '第二段包含足够文字供正常换行，内容仍来自合成回归输入。',
                    '第三段说明本检查不替代实际分页或逐页视觉验收。'],
         '关键词：合成；摘要；布局', '宋体'),
        ('ABSTRACT', ['This synthetic abstract checks Word paragraph layout and retains A1 and kW.',
                      'The second paragraph contains synthetic input for ordinary line wrapping.',
                      'The third paragraph does not establish pagination or visual acceptance.'],
         'KEYWORDS: synthetic; abstract; layout', 'Times New Roman'),
    ):
        document.add_paragraph(title)
        for index, text in enumerate([*texts, keywords]):
            paragraph = document.add_paragraph()
            midpoint = len(text) // 2
            paragraph.add_run(text[:midpoint])
            paragraph.add_run(text[midpoint:])
            if index % 2 == 0:
                snap = OxmlElement('w:snapToGrid')
                snap.set(qn('w:val'), '1')
                paragraph._p.get_or_add_pPr().append(snap)
            expected[text] = east_asia
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_paragraph('目录')
    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_paragraph('1 合成正文', style='Heading 1')
    for section in document.sections:
        for old_grid in list(section._sectPr.findall(qn('w:docGrid'))):
            section._sectPr.remove(old_grid)
        grid = OxmlElement('w:docGrid')
        grid.set(qn('w:type'), 'lines')
        grid.set(qn('w:linePitch'), '312')
        section._sectPr.append(grid)
    return document, expected


class ParagraphGridSpacing(unittest.TestCase):
    def check_spacing(self, formatter, text):
        document = Document()
        grid = OxmlElement('w:docGrid')
        grid.set(qn('w:type'), 'lines')
        grid.set(qn('w:linePitch'), '312')
        document.sections[0]._sectPr.append(grid)
        paragraph = document.add_paragraph(text)
        formatter(paragraph)
        snap = paragraph._p.pPr.find(qn('w:snapToGrid'))
        self.assertIsNotNone(snap, 'an absent override inherits the template grid')
        self.assertEqual(snap.get(qn('w:val')), '0')
        self.assertEqual(paragraph.paragraph_format.line_spacing, 1.5)
        self.assertEqual(paragraph.text, text)
        self.assertEqual(paragraph.runs[0].font.size.pt, 12)
        self.assertEqual(grid.get(qn('w:linePitch')), '312', 'retain the source template grid')

    def test_wrapped_body_keeps_declared_spacing(self):
        self.check_spacing(_format_body, 'REPORT-REVISION-20261003：本次测试使用独立目录与受控材料，不构成生产结论。')

    def test_wrapped_reference_keeps_declared_spacing(self):
        self.check_spacing(_format_reference, '[1] 测试组. 固定输入均值计算组件测试记录[R]. 2026.')

    def test_bilingual_abstracts_and_keywords_override_grid_after_docx_readback(self):
        document, expected = abstract_fixture()

        def section_layout(document):
            return [{tag: dict(section._sectPr.find(qn('w:' + tag)).attrib)
                     for tag in ('pgSz', 'pgMar', 'docGrid')} for section in document.sections]

        original_layout = section_layout(document)
        _configure_front_matter(document)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'synthetic-abstract.docx'
            document.save(path)
            readback = Document(path)
        self.assertEqual(section_layout(readback), original_layout)
        paragraphs = {paragraph.text: paragraph for paragraph in readback.paragraphs}
        for text, east_asia in expected.items():
            with self.subTest(text=text):
                self.assertIn(text, paragraphs, 'abstract and keyword text survives formatting')
                paragraph = paragraphs[text]
                snaps = paragraph._p.pPr.findall(qn('w:snapToGrid'))
                self.assertEqual(len(snaps), 1)
                self.assertEqual(snaps[0].get(qn('w:val')), '0')
                spacing = paragraph._p.pPr.find(qn('w:spacing'))
                self.assertEqual(spacing.get(qn('w:line')), '360')
                self.assertEqual(spacing.get(qn('w:lineRule')), 'auto')
                self.assertEqual(spacing.get(qn('w:before')), '0')
                self.assertEqual(spacing.get(qn('w:after')), '0')
                self.assertEqual(len(paragraph.runs), 2, 'retain both source text runs')
                for run in paragraph.runs:
                    self.assertEqual(run.font.size.pt, 12)
                    fonts = run._r.rPr.find(qn('w:rFonts'))
                    self.assertEqual(fonts.get(qn('w:eastAsia')), east_asia)
                    self.assertEqual(fonts.get(qn('w:ascii')), 'Times New Roman')
                    self.assertEqual(fonts.get(qn('w:hAnsi')), 'Times New Roman')


if __name__ == '__main__':
    unittest.main(verbosity=2)
