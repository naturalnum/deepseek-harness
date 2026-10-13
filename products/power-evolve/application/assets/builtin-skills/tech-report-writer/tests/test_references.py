"""Dataset references through Word JSON validation and DOCX paragraph readback.

The DOCX fixture is intentionally minimal: only reference errors are asserted,
so these tests do not claim that its unrelated report layout passes validation.
"""
import html
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from docx import Document
from main import _validate_references_data
from report_validator import _gbt7714_reference_error, validate_report

URL = "https://archive.ics.uci.edu/dataset/321/electricityloaddiagrams20112014"
ENTRY = ("TRINDADE A. ElectricityLoadDiagrams20112014[DS/OL]. "
         "UCI Machine Learning Repository, 2015[2026-10-03]. " + URL + ". DOI:10.24432/C58C86.")


def references(entries):
    return {"references_html": "".join("<p>" + html.escape(value) + "</p>" for value in entries)}


class DatasetReferences(unittest.TestCase):
    def docx_reference_errors(self, entries):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "references.docx"
            document = Document()
            document.add_paragraph("参考文献")
            for entry in entries:
                document.add_paragraph(entry)
            document.save(path)
            result = validate_report(str(path))
            return [error for error in result["errors"] if "参考文献" in error]

    def test_actual_qa_dataset_word_cli(self):
        with tempfile.TemporaryDirectory() as directory:
            data = references(["[1] " + ENTRY])
            (Path(directory) / "references.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
            result = subprocess.run([sys.executable, str(ROOT / "scripts/main.py"), "validate-references"],
                                    capture_output=True, text=True, timeout=10,
                                    env=dict(os.environ, TECH_REPORT_OUTPUT_DIR=directory, PYTHONDONTWRITEBYTECODE="1"))
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_actual_qa_dataset_docx_readback(self):
        self.assertEqual(self.docx_reference_errors(["[1] " + ENTRY]), [])

    def test_bare_and_online_dataset_markers(self):
        for marker in ("DS", "DS/OL", "DB", "DB/OL", "R", "J"):
            with self.subTest(marker=marker):
                entry = ENTRY.replace("DS/OL", marker)
                if "/" not in marker:
                    entry = entry.replace("[2026-10-03]", "").replace(URL + ". ", "")
                self.assertTrue(_validate_references_data(references(["[1] " + entry]))[0])
                self.assertIsNone(_gbt7714_reference_error(entry))

    def test_original_missing_url_and_other_entry_failures(self):
        cases = {
            "original QA DOI without URL": ENTRY.replace(URL + ". ", ""),
            "missing date": ENTRY.replace("[2026-10-03]", ""),
            "missing final period": ENTRY[:-1],
            "unsupported dataset carrier": ENTRY.replace("DS/OL", "DS/DK"),
            "unsupported document marker": ENTRY.replace("DS/OL", "UNKNOWN/OL"),
        }
        for name, entry in cases.items():
            with self.subTest(case=name):
                self.assertFalse(_validate_references_data(references(["[1] " + entry]))[0])
                self.assertIsNotNone(_gbt7714_reference_error(entry))
                self.assertTrue(self.docx_reference_errors(["[1] " + entry]))

    def test_final_sequence_numbers_remain_continuous(self):
        for entries in (["[4] " + ENTRY], ["[1] " + ENTRY, "[3] " + ENTRY], [ENTRY]):
            with self.subTest(entries=entries):
                self.assertFalse(_validate_references_data(references(entries))[0])
                self.assertTrue(self.docx_reference_errors(entries))


if __name__ == "__main__":
    unittest.main(verbosity=2)
