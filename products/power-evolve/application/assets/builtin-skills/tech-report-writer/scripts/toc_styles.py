"""TOC / figure-list / table-list style and field helpers.

The main table of contents, list of figures and list of tables intentionally share
one visual system: same TOC entry styles, fonts, size, line spacing, tab leaders
and page-number alignment.  Figure/table entries are collected through hidden TC
fields so the implementation does not depend on Word SEQ captions.
"""

from __future__ import annotations

import copy
import re

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


# A real caption must contain an explicit separator between the number and the
# title.  This deliberately excludes explanatory prose such as
# ``图1-1应表达……`` / ``表2-1给出……`` which used to be mistaken for catalog
# entries by the permissive ``\s*`` pattern.
FIGURE_CAPTION_RE = re.compile(
    r"^图\s*(?P<number>[A-Z0-9]+-\d+)[ \t\u3000]+(?P<title>\S.*)$"
)
TABLE_CAPTION_RE = re.compile(
    r"^表\s*(?P<number>[A-Z0-9]+-\d+)[ \t\u3000]+(?P<title>\S.*)$"
)

_EXPLANATION_TITLE_PREFIX_RE = re.compile(
    r"^(?:应表达|应展示|应给出|给出|展示|显示|表明|说明|反映|可见|"
    r"所示|用于表达|用于展示|用于说明|用于给出)"
)


def configure_toc_styles(doc) -> None:
    """Normalize Word/WPS TOC styles for all three catalog fields."""
    harmful_tags = ("w:b", "w:bCs", "w:i", "w:iCs", "w:smallCaps", "w:caps")

    try:
        toc_style_names = []
        for style in doc.styles:
            name = style.name or ""
            if name.startswith("toc ") or name.startswith("目录 "):
                toc_style_names.append(name)

        for style_name in toc_style_names:
            try:
                style = doc.styles[style_name]
            except KeyError:
                continue

            font = style.font
            font.size = Pt(12)
            font.bold = False
            font.italic = False
            font.name = None
            style.paragraph_format.line_spacing = 1.0

            style_element = style.element
            run_properties = style_element.find(qn("w:rPr"))
            if run_properties is None:
                run_properties = OxmlElement("w:rPr")
                name_element = style_element.find(qn("w:name"))
                if name_element is not None:
                    style_element.insert(
                        list(style_element).index(name_element), run_properties
                    )
                else:
                    style_element.append(run_properties)

            old_fonts = run_properties.find(qn("w:rFonts"))
            if old_fonts is not None:
                run_properties.remove(old_fonts)

            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:eastAsia"), "宋体")
            fonts.set(qn("w:ascii"), "Times New Roman")
            fonts.set(qn("w:hAnsi"), "Times New Roman")
            run_properties.append(fonts)

            for tag in harmful_tags:
                for element in run_properties.findall(qn(tag)):
                    run_properties.remove(element)

            size = run_properties.find(qn("w:sz"))
            size_complex = run_properties.find(qn("w:szCs"))
            if size is not None and size_complex is not None:
                size_complex.set(qn("w:val"), size.get(qn("w:val")))
    except Exception:
        pass


def _instruction_text(paragraph) -> str:
    return "".join(
        element.text or "" for element in paragraph._p.iter(qn("w:instrText"))
    )


def _field_run(kind: str, text: str | None = None, hidden: bool = False):
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    if hidden:
        vanish = OxmlElement("w:vanish")
        rpr.append(vanish)
    run.append(rpr)
    if kind == "begin":
        node = OxmlElement("w:fldChar")
        node.set(qn("w:fldCharType"), "begin")
        node.set(qn("w:dirty"), "true")
    elif kind == "separate":
        node = OxmlElement("w:fldChar")
        node.set(qn("w:fldCharType"), "separate")
    elif kind == "end":
        node = OxmlElement("w:fldChar")
        node.set(qn("w:fldCharType"), "end")
    elif kind == "instr":
        node = OxmlElement("w:instrText")
        node.set(qn("xml:space"), "preserve")
        node.text = text or ""
    else:
        node = OxmlElement("w:t")
        node.set(qn("xml:space"), "preserve")
        node.text = text or ""
    run.append(node)
    return run


def _toc_field_paragraph(identifier: str, prompt: str):
    """Create a custom-TC TOC field; entries use normal TOC 1 styles."""
    p = OxmlElement("w:p")
    p.append(_field_run("begin"))
    p.append(_field_run("instr", f' TOC \\f {identifier} \\h \\z '))
    p.append(_field_run("separate"))
    p.append(_field_run("text", prompt))
    p.append(_field_run("end"))
    return p


def _move_section_properties(source_p, target_p) -> None:
    """Move a paragraph-level section break so catalogs stay in front matter."""
    src_ppr = source_p.find(qn("w:pPr"))
    if src_ppr is None:
        return
    sect = src_ppr.find(qn("w:sectPr"))
    if sect is None:
        return
    src_ppr.remove(sect)
    tgt_ppr = target_p.find(qn("w:pPr"))
    if tgt_ppr is None:
        tgt_ppr = OxmlElement("w:pPr")
        target_p.insert(0, tgt_ppr)
    old = tgt_ppr.find(qn("w:sectPr"))
    if old is not None:
        tgt_ppr.remove(old)
    tgt_ppr.append(copy.deepcopy(sect))


def _find_toc_paragraph(doc, marker: str | None = None):
    fallback = None
    for paragraph in doc.paragraphs:
        instruction = _instruction_text(paragraph)
        if "TOC" not in instruction:
            continue
        if marker is None:
            if re.search(r"\\f\s+[FT](?:\s|$)", instruction):
                continue
            if fallback is None:
                fallback = paragraph
            if '\\o "1-3"' in instruction or "\\o '1-3'" in instruction:
                return paragraph
        elif re.search(rf"\\f\s+{re.escape(marker)}(?:\s|$)", instruction):
            return paragraph
    return fallback if marker is None else None


def _remove_legacy_catalog_titles(doc) -> None:
    """Remove visible legacy labels/separator paragraphs from older releases.

    v1.1.22 keeps the figure/table TOC fields as continuations of the main TOC.
    Therefore visible ``图目录`` / ``表目录`` heading paragraphs are deleted.
    Empty paragraphs that existed only as separators immediately around those
    legacy labels are removed as well.
    """
    paragraphs = list(doc.paragraphs)
    remove_nodes = []
    for i, paragraph in enumerate(paragraphs):
        text = re.sub(r"\s+", "", paragraph.text or "")
        if text not in {"图目录", "表目录"}:
            continue
        remove_nodes.append(paragraph._p)
        # Older versions could leave an empty visual spacer next to the title.
        for j in (i - 1, i + 1):
            if 0 <= j < len(paragraphs):
                neighbor = paragraphs[j]
                if not (neighbor.text or "").strip() and not _instruction_text(neighbor):
                    remove_nodes.append(neighbor._p)
    seen = set()
    for node in remove_nodes:
        if id(node) in seen:
            continue
        seen.add(id(node))
        parent = node.getparent()
        if parent is not None:
            parent.remove(node)


def _clear_legacy_catalog_prompt(paragraph) -> None:
    """Remove only the old visible update prompt from a custom catalog field."""
    for text_node in paragraph._p.iter(qn("w:t")):
        text = text_node.text or ""
        if "更新域" in text and ("WPS" in text or "Word" in text):
            text_node.text = ""


def ensure_figure_table_catalogs(doc) -> None:
    """Append figure/table entries directly after the main TOC, without labels.

    The three dynamic TOC fields remain separate fields so Word/WPS can update
    them independently, but no visible ``图目录`` / ``表目录`` heading and no
    dedicated spacer paragraph is inserted.  Figure/table entries inherit the
    same TOC styles, fonts, size, single line spacing, tab leaders and page
    number alignment as the main TOC.

    If the main TOC paragraph owns the front-matter/body section break, move
    that section break to the final table-list field so all catalog entries stay
    in the Roman-numbered front matter.
    """
    configure_toc_styles(doc)
    _remove_legacy_catalog_titles(doc)
    main = _find_toc_paragraph(doc, None)
    if main is None:
        return

    fig_field = _find_toc_paragraph(doc, "F")
    table_field = _find_toc_paragraph(doc, "T")
    if fig_field is not None and table_field is not None:
        _clear_legacy_catalog_prompt(fig_field)
        _clear_legacy_catalog_prompt(table_field)
        for paragraph in (main, fig_field, table_field):
            for field in paragraph._p.iter(qn("w:fldChar")):
                if field.get(qn("w:fldCharType")) == "begin":
                    field.set(qn("w:dirty"), "true")
        _move_section_properties(main._p, table_field._p)
        return

    anchor = main._p
    # Empty result text is intentional: before fields are updated there is no
    # visible catalog label/prompt; after update, entries begin directly after
    # the preceding TOC entries.
    fig_p = _toc_field_paragraph("F", "")
    table_p = _toc_field_paragraph("T", "")

    for element in (fig_p, table_p):
        anchor.addnext(element)
        anchor = element

    _move_section_properties(main._p, table_p)

def _escape_tc_text(text: str) -> str:
    return text.replace('\\', '\\\\').replace('"', '\\"')


def _has_tc_identifier(paragraph, identifier: str) -> bool:
    instruction = _instruction_text(paragraph)
    return bool(re.search(rf"\bTC\b.*\\f\s+{re.escape(identifier)}(?:\s|$)", instruction))


def _remove_tc_identifier(paragraph, identifier: str) -> None:
    """Remove complete TC fields for one catalog identifier only.

    Older releases could leave a stale F/T field on an explanatory paragraph.
    The cleanup is intentionally range-based so unrelated Word fields in the
    same paragraph are preserved.
    """
    children = list(paragraph._p)
    remove_indexes: set[int] = set()
    for index, child in enumerate(children):
        instruction = "".join(
            element.text or "" for element in child.iter(qn("w:instrText"))
        )
        if not re.search(
            rf"\bTC\b.*\\f\s+{re.escape(identifier)}(?:\s|$)", instruction
        ):
            continue
        start = index
        while start >= 0:
            fld = children[start].find(qn("w:fldChar"))
            if fld is not None and fld.get(qn("w:fldCharType")) == "begin":
                break
            start -= 1
        end = index
        while end < len(children):
            fld = children[end].find(qn("w:fldChar"))
            if fld is not None and fld.get(qn("w:fldCharType")) == "end":
                break
            end += 1
        if start >= 0 and end < len(children):
            remove_indexes.update(range(start, end + 1))
    for index in sorted(remove_indexes, reverse=True):
        child = children[index]
        if child.getparent() is paragraph._p:
            paragraph._p.remove(child)


def _has_exact_tc(paragraph, text: str, identifier: str) -> bool:
    safe_text = re.escape(text.strip())
    instruction = _instruction_text(paragraph)
    return bool(
        re.search(rf'\bTC\s+"{safe_text}"', instruction)
        and re.search(rf"\\f\s+{re.escape(identifier)}(?:\s|$)", instruction)
        and re.search(r'\\l\s+"1"', instruction)
    )


def _append_hidden_tc(paragraph, text: str, identifier: str) -> None:
    # Keep a correct marker, but replace stale text or a malformed marker.
    if _has_exact_tc(paragraph, text, identifier):
        return
    if _has_tc_identifier(paragraph, identifier):
        _remove_tc_identifier(paragraph, identifier)
    safe = _escape_tc_text(text.strip())
    paragraph._p.append(_field_run("begin", hidden=True))
    paragraph._p.append(_field_run("instr", f' TC "{safe}" \\f {identifier} \\l "1" ', hidden=True))
    paragraph._p.append(_field_run("separate", hidden=True))
    paragraph._p.append(_field_run("end", hidden=True))


def _paragraph_has_embedded_figure(paragraph) -> bool:
    return any(
        paragraph._p.find(f".//{qn(tag)}") is not None
        for tag in ("w:drawing", "w:pict", "w:object")
    )


def _previous_paragraph_has_figure(paragraph) -> bool:
    node = paragraph._p.getprevious()
    while node is not None:
        if node.tag == qn("w:tbl"):
            return False
        if node.tag == qn("w:p"):
            if any(
                node.find(f".//{qn(tag)}") is not None
                for tag in ("w:drawing", "w:pict", "w:object")
            ):
                return True
            visible = "".join(t.text or "" for t in node.iter(qn("w:t"))).strip()
            if visible:
                return False
        node = node.getprevious()
    return False


def _next_block_is_table(paragraph) -> bool:
    node = paragraph._p.getnext()
    while node is not None:
        if node.tag == qn("w:tbl"):
            return True
        if node.tag == qn("w:p"):
            visible = "".join(t.text or "" for t in node.iter(qn("w:t"))).strip()
            if visible:
                return False
        node = node.getnext()
    return False


def _is_true_caption(paragraph, kind: str) -> bool:
    text = re.sub(r"\s+", " ", paragraph.text.strip())
    match = (FIGURE_CAPTION_RE if kind == "figure" else TABLE_CAPTION_RE).match(text)
    if match is None:
        return False
    # Continued-table labels are layout continuations, not new catalog items.
    if kind == "table" and re.match(r"^[（(]续[)）]", match.group("title")):
        return False
    if _EXPLANATION_TITLE_PREFIX_RE.match(match.group("title").lstrip("：: ")):
        return False

    style_name = (getattr(paragraph.style, "name", "") or "").lower()
    if style_name == "caption":
        return True
    # Generated captions are centered. Existing reports are also accepted when
    # the expected adjacent object proves the paragraph is a real caption.
    try:
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        if paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER:
            return True
    except Exception:
        pass
    if kind == "figure":
        return _previous_paragraph_has_figure(paragraph)
    return _next_block_is_table(paragraph)


def ensure_abstract_catalog_marks(doc) -> dict[str, bool]:
    """Ensure 内容摘要 / ABSTRACT each expose one hidden main-TOC TC mark."""
    status = {"内容摘要": False, "ABSTRACT": False}
    for paragraph in doc.paragraphs:
        label = paragraph.text.strip()
        if label not in status:
            continue
        _append_hidden_tc(paragraph, label, "C")
        status[label] = True
    return status


def mark_figure_table_captions(doc) -> dict[str, int]:
    """Mark only real caption paragraphs for the separate catalogs.

    Stale F/T marks on ordinary prose are removed on every pass. This makes the
    function safe both for newly generated reports and for re-formatting DOCX
    files created by older releases.
    """
    counts = {"figures": 0, "tables": 0}
    for paragraph in doc.paragraphs:
        text = re.sub(r"\s+", " ", paragraph.text.strip())
        if _is_true_caption(paragraph, "figure"):
            _remove_tc_identifier(paragraph, "T")
            _append_hidden_tc(paragraph, text, "F")
            counts["figures"] += 1
        elif _is_true_caption(paragraph, "table"):
            _remove_tc_identifier(paragraph, "F")
            _append_hidden_tc(paragraph, text, "T")
            counts["tables"] += 1
        else:
            _remove_tc_identifier(paragraph, "F")
            _remove_tc_identifier(paragraph, "T")
    return counts


def catalog_field_status(doc) -> dict[str, bool]:
    """Small validator helper used by report_validator and tests."""
    return {
        "main_toc": _find_toc_paragraph(doc, None) is not None,
        "figure_catalog": _find_toc_paragraph(doc, "F") is not None,
        "table_catalog": _find_toc_paragraph(doc, "T") is not None,
    }
