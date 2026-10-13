"""
图片处理模块。

判断<img>标签的图片来源（Base64 / 本地文件 / 网络URL），
获取图片字节流并插入到Word文档中。

尊重 width 和 height 属性，保持宽高比。

所有函数不抛出异常，加载失败时返回警告信息。
"""

import base64
import os
import os.path
import re
from io import BytesIO
from typing import List, Optional, Set, Tuple

from docx.shape import InlineShape
from docx.shared import Inches
from PIL import Image as PILImage

# ---------------------------------------------------------------------------
# 图片下载超时设置（秒）
# ---------------------------------------------------------------------------
_IMAGE_DOWNLOAD_TIMEOUT = 30

# ---------------------------------------------------------------------------
# 图片字节数上限（10 MB），防止恶意超大图片导致内存溢出
# ---------------------------------------------------------------------------
MAX_IMAGE_BYTES = 10 * 1024 * 1024  # 10 MB
MAX_BASE64_STRING_LENGTH = 14 * 1024 * 1024  # 14 MB

# ---------------------------------------------------------------------------
# 默认图片尺寸限制（英寸）
# ---------------------------------------------------------------------------
DEFAULT_MAX_WIDTH_INCHES = 6.0
DEFAULT_MAX_HEIGHT_INCHES = 8.0

# ---------------------------------------------------------------------------
# 本地图片路径白名单
# ---------------------------------------------------------------------------
# None 表示不限制（向后兼容）；设为集合后，只允许解析后的绝对路径
# 位于集合中任一目录之内。由 generate() 入口统一设置。
_allowed_image_dirs: Optional[Set[str]] = None


def set_allowed_image_dirs(dirs: Optional[List[str]]) -> None:
    """
    设置本地图片路径白名单。

    调用 generate() 或 generate_from_template() 时自动调用。传入 None 表示不限制路径
    （向后兼容），传入目录列表则后续本地图片只能从这些目录下加载。

    Args:
        dirs: 允许的目录路径列表，或 None 表示不限制。
    """
    global _allowed_image_dirs
    if dirs is None:
        _allowed_image_dirs = None
        return
    # 规范化所有路径为绝对路径
    _allowed_image_dirs = set()
    for d in dirs:
        try:
            abs_d = os.path.abspath(os.path.realpath(d))
            # 确保以分隔符结尾，便于前缀匹配
            _allowed_image_dirs.add(abs_d.rstrip(os.sep) + os.sep)
        except Exception:
            pass  # 无效路径跳过


def _get_image_bytes(src: str) -> Tuple[Optional[BytesIO], Optional[str]]:
    """
    根据 src 属性获取图片字节流。

    支持三种来源：
    1. Base64编码：以 data:image/...;base64, 开头
    2. 本地文件路径：不以 http 或 data: 开头
    3. 网络URL：以 http:// 或 https:// 开头

    Args:
        src: <img> 标签的 src 属性值。

    Returns:
        (BytesIO对象或None, 警告字符串或None)
    """
    if not src or not isinstance(src, str):
        return None, "[图片] src为空"

    src = src.strip()

    if not src:
        return None, "[图片] src为空"

    try:
        # --- 来源1：Base64编码 ---
        if src.startswith("data:"):
            return _load_base64_image(src)

        # --- 来源2：网络URL ---
        if src.startswith("http://") or src.startswith("https://"):
            return _load_url_image(src)

        # --- 来源3：本地文件路径 ---
        return _load_local_image(src)

    except Exception as e:
        return None, f"[图片] 无法加载: {src} ({e})"


def _load_base64_image(src: str) -> Tuple[Optional[BytesIO], Optional[str]]:
    """
    解析并解码 Base64 编码的图片。

    格式：data:image/png;base64,iVBORw0KGgo...

    包含大小限制检查，防止恶意超大 Base64 字符串导致内存溢出。

    Args:
        src: Base64编码的图片数据URL。

    Returns:
        (BytesIO对象或None, 警告字符串或None)
    """
    # --- 大小限制（步骤1）：检查原始字符串长度 ---
    if len(src) > MAX_BASE64_STRING_LENGTH:
        return None, (
            f"[图片] Base64数据过大（{len(src)} 字符，上限 {MAX_BASE64_STRING_LENGTH}），已拒绝处理以防内存溢出"
        )

    # 匹配 data:image/xxx;base64,<data>
    match = re.match(r"data:image/\w+;base64,(.+)", src, re.IGNORECASE)
    if not match:
        return None, "[图片] 无效的Base64数据格式"

    data_str = match.group(1)

    # --- 大小限制（步骤2）：解码前估算数据大小（Base64→原始 约 3/4） ---
    estimated_size = (len(data_str) * 3) // 4
    if estimated_size > MAX_IMAGE_BYTES:
        return None, (
            f"[图片] Base64解码后数据量过大（约 {estimated_size} 字节，"
            f"上限 {MAX_IMAGE_BYTES} 字节），已拒绝处理以防内存溢出"
        )

    try:
        image_data = base64.b64decode(data_str)

        # --- 大小限制（步骤3）：解码后最终检查 ---
        if len(image_data) > MAX_IMAGE_BYTES:
            return None, (
                f"[图片] 解码后图片过大（{len(image_data)} 字节，上限 {MAX_IMAGE_BYTES} 字节），已拒绝处理以防内存溢出"
            )

        return BytesIO(image_data), None
    except Exception as e:
        return None, f"[图片] Base64解码失败: {e}"


def _load_url_image(url: str) -> Tuple[Optional[BytesIO], Optional[str]]:
    """
    从网络URL下载图片。

    使用项目统一的 request_client 发起请求，具备自动重试能力。
    下载后检查数据量，超过 MAX_IMAGE_BYTES 时拒绝处理。

    Args:
        url: 图片的网络地址。

    Returns:
        (BytesIO对象或None, 警告字符串或None)
    """
    import requests

    from engine.request_client import request_get

    try:
        response = request_get(
            url,
            timeout=_IMAGE_DOWNLOAD_TIMEOUT,
            headers={"User-Agent": "html-to-word/1.0"},
        )
        if response.status_code == 200:
            content = response.content
            if len(content) > MAX_IMAGE_BYTES:
                return None, (
                    f"[图片] 下载的图片过大（{len(content)} 字节，上限 {MAX_IMAGE_BYTES} 字节），已拒绝处理以防内存溢出"
                )
            return BytesIO(content), None
        else:
            return None, f"[图片] 下载失败，HTTP状态码: {response.status_code}"
    except requests.exceptions.Timeout:
        return None, f"[图片] 下载超时: {url}"
    except requests.exceptions.ConnectionError:
        return None, f"[图片] 连接失败: {url}"
    except requests.exceptions.RequestException as e:
        return None, f"[图片] 网络请求异常: {url} ({e})"
    except Exception as e:
        return None, f"[图片] 无法加载: {url} ({e})"


def _load_local_image(path: str) -> Tuple[Optional[BytesIO], Optional[str]]:
    """
    从本地文件系统读取图片。

    安全检查：若已设置 allowed_image_dirs，则解析路径的绝对路径后
    检查是否位于允许的目录内，拒绝目录遍历攻击（如 ../../etc/passwd）。
    未设置白名单时保持向后兼容，不做限制。

    Args:
        path: 本地图片文件路径。

    Returns:
        (BytesIO对象或None, 警告字符串或None)
    """
    try:
        # ---- 路径安全检查：防目录遍历 ----
        if _allowed_image_dirs is not None:
            real_path = os.path.realpath(path)
            abs_real = os.path.abspath(real_path)
            allowed = any(
                abs_real.startswith(allowed_dir) or (abs_real + os.sep).startswith(allowed_dir)
                for allowed_dir in _allowed_image_dirs
            )
            if not allowed:
                return None, (f"[图片] 路径不在允许目录内: {path}，已拒绝以防止路径遍历攻击")

        # ---- 大小限制：预检查文件大小 ----
        file_size = os.path.getsize(path)
        if file_size > MAX_IMAGE_BYTES:
            return None, (f"[图片] 文件过大（{file_size} 字节，上限 {MAX_IMAGE_BYTES} 字节），已拒绝处理以防内存溢出")

        with open(path, "rb") as f:
            data = f.read()
        return BytesIO(data), None
    except FileNotFoundError:
        return None, f"[图片] 文件不存在: {path}"
    except PermissionError:
        return None, f"[图片] 无权限读取文件: {path}"
    except Exception as e:
        return None, f"[图片] 无法加载: {path} ({e})"


def handle_image(
    doc,
    img_tag,
    max_width_inches: float = DEFAULT_MAX_WIDTH_INCHES,
    max_height_inches: float = DEFAULT_MAX_HEIGHT_INCHES,
) -> Tuple[Optional[InlineShape], List[str]]:
    """
    处理<img>标签，获取图片并插入到文档中。

    根据 width 和 height 属性计算缩放比例，保持宽高比。
    使用 PIL 读取图片的实际尺寸作为计算基准。

    Args:
        doc: python-docx 的 Document 对象。
        img_tag: BeautifulSoup 的 <img> Tag 对象。
        max_width_inches: 图片最大宽度（英寸）。
        max_height_inches: 图片最大高度（英寸）。

    Returns:
        (插入的图片对象或None, 警告字符串列表)
    """
    warnings: List[str] = []

    src = img_tag.get("src", "") if hasattr(img_tag, "get") else ""
    image_bytes, warning = _get_image_bytes(src)

    if warning:
        warnings.append(warning)

    if image_bytes is None:
        return None, warnings

    try:
        # 获取 width 和 height 属性
        attr_width = img_tag.get("width", None) if hasattr(img_tag, "get") else None
        attr_height = img_tag.get("height", None) if hasattr(img_tag, "get") else None

        # 使用 PIL 获取图片实际尺寸
        image_bytes.seek(0)
        pil_image = PILImage.open(image_bytes)
        actual_width_px, actual_height_px = pil_image.size

        # 确定目标宽度和高度（像素单位）
        target_width_px = _parse_pixel_value(attr_width) if attr_width else actual_width_px
        target_height_px = _parse_pixel_value(attr_height) if attr_height else actual_height_px

        # 如果只指定了一个维度，按比例计算另一个
        if attr_width and not attr_height:
            ratio = target_width_px / actual_width_px
            target_height_px = actual_height_px * ratio
        elif attr_height and not attr_width:
            ratio = target_height_px / actual_height_px
            target_width_px = actual_width_px * ratio

        # 转换为英寸（1英寸 = 96px，Word标准）
        width_inches = target_width_px / 96.0
        height_inches = target_height_px / 96.0

        # 限制最大尺寸
        if width_inches > max_width_inches:
            ratio = max_width_inches / width_inches
            width_inches = max_width_inches
            height_inches *= ratio
        if height_inches > max_height_inches:
            ratio = max_height_inches / height_inches
            height_inches = max_height_inches
            width_inches *= ratio

        # 确保最小尺寸
        width_inches = max(width_inches, 0.1)
        height_inches = max(height_inches, 0.1)

        # 插入图片
        image_bytes.seek(0)
        picture = doc.add_picture(
            image_bytes,
            width=Inches(width_inches),
            height=Inches(height_inches),
        )

        # 关闭PIL图片
        pil_image.close()

        return picture, warnings

    except Exception as e:
        warnings.append(f"[图片] 处理图片时发生错误: ({e})")
        return None, warnings


def handle_image_to_paragraph(
    para,
    img_tag,
    max_width_inches: float = DEFAULT_MAX_WIDTH_INCHES,
    max_height_inches: float = DEFAULT_MAX_HEIGHT_INCHES,
) -> List[str]:
    """
    处理<img>标签，将图片插入到指定段落中。

    与 handle_image 不同，此函数不创建新段落，而是将图片作为 run
    插入到已有段落中，确保图片位于占位符的正确位置。

    Args:
        para: 目标段落（python-docx Paragraph 对象）。
        img_tag: BeautifulSoup 的 <img> Tag 对象。
        max_width_inches: 图片最大宽度（英寸）。
        max_height_inches: 图片最大高度（英寸）。

    Returns:
        警告字符串列表。
    """
    warnings: List[str] = []

    src = img_tag.get("src", "") if hasattr(img_tag, "get") else ""
    image_bytes, warning = _get_image_bytes(src)

    if warning:
        warnings.append(warning)

    if image_bytes is None:
        return warnings

    try:
        attr_width = img_tag.get("width", None) if hasattr(img_tag, "get") else None
        attr_height = img_tag.get("height", None) if hasattr(img_tag, "get") else None

        image_bytes.seek(0)
        pil_image = PILImage.open(image_bytes)
        actual_width_px, actual_height_px = pil_image.size

        target_width_px = _parse_pixel_value(attr_width) if attr_width else actual_width_px
        target_height_px = _parse_pixel_value(attr_height) if attr_height else actual_height_px

        if attr_width and not attr_height:
            ratio = target_width_px / actual_width_px
            target_height_px = actual_height_px * ratio
        elif attr_height and not attr_width:
            ratio = target_height_px / actual_height_px
            target_width_px = actual_width_px * ratio

        width_inches = target_width_px / 96.0
        height_inches = target_height_px / 96.0

        if width_inches > max_width_inches:
            ratio = max_width_inches / width_inches
            width_inches = max_width_inches
            height_inches *= ratio
        if height_inches > max_height_inches:
            ratio = max_height_inches / height_inches
            height_inches = max_height_inches
            width_inches *= ratio

        width_inches = max(width_inches, 0.1)
        height_inches = max(height_inches, 0.1)

        image_bytes.seek(0)
        run = para.add_run()
        run.add_picture(image_bytes, width=Inches(width_inches), height=Inches(height_inches))

        pil_image.close()

    except Exception as e:
        warnings.append(f"[图片] 处理图片时发生错误: ({e})")

    return warnings


def _parse_pixel_value(value) -> int:
    """
    解析像素值字符串。

    支持格式：'100', '100px'

    Args:
        value: 像素值字符串或数字。

    Returns:
        整数值像素。
    """
    if isinstance(value, (int, float)):
        return int(value)
    if isinstance(value, str):
        value = value.strip().lower().rstrip("px").strip()
        try:
            return int(float(value))
        except (ValueError, TypeError):
            return 0
    return 0
