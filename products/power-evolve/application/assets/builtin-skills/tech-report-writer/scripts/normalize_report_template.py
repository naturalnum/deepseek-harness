#!/usr/bin/env python3
"""将 report.docx 规范化为中国电科院技术报告模板。"""

import os
import re
import sys
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from config import (
    COVER_BOTTOM_FONT_SIZE_PT,
    COVER_HEITI_FONT,
    COVER_METADATA_FONT_SIZE_PT,
    COVER_PROJECT_SOURCE_FONT_SIZE_PT,
    COVER_SERIES_FONT_SIZE_PT,
    COVER_TITLE_CN_FONT_SIZE_PT,
    FOOTER_DISTANCE_MM,
    HEADER_DISTANCE_MM,
    PAGE_HEIGHT_MM,
    PAGE_MARGIN_BOTTOM_MM,
    PAGE_MARGIN_LEFT_MM,
    PAGE_MARGIN_RIGHT_MM,
    PAGE_MARGIN_TOP_MM,
    PAGE_WIDTH_MM,
    REPORT_WRITING_UNIT,
)
from cover_bottom_layout import normalize_cover_bottom_template
from cover_layout import calculate_cover_metadata_indents_twips


def _normalize_cover_bottom_block(doc):
    normalize_cover_bottom_template(doc)


def _align_cover_metadata_block(paragraphs):
    """三行左对齐，并让最长一行的最右字符贴齐右页边距内侧。"""
    paragraphs = list(paragraphs)
    left_twips, right_twips = calculate_cover_metadata_indents_twips(
        paragraph.text for paragraph in paragraphs
    )
    for paragraph in paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_pr = paragraph._p.get_or_add_pPr()

        for tabs in list(p_pr.findall(qn("w:tabs"))):
            p_pr.remove(tabs)

        indent = p_pr.find(qn("w:ind"))
        if indent is None:
            indent = OxmlElement("w:ind")
            p_pr.append(indent)
        for name in (
            "left",
            "leftChars",
            "right",
            "rightChars",
            "firstLine",
            "firstLineChars",
            "hanging",
            "hangingChars",
        ):
            indent.attrib.pop(qn(f"w:{name}"), None)
        indent.set(qn("w:left"), str(left_twips))
        indent.set(qn("w:right"), str(right_twips))


def _find_paragraph(doc, text):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"模板中未找到段落: {text}")


def _set_text_preserving_first_run(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def _remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)


def _set_two_character_first_line_indent(paragraph):
    """设置首行缩进 2 字符，并清除模板遗留的左右缩进。"""
    p_pr = paragraph._p.get_or_add_pPr()
    indent = p_pr.find(qn("w:ind"))
    if indent is None:
        indent = OxmlElement("w:ind")
        p_pr.append(indent)

    for name in ("left", "leftChars", "right", "rightChars", "hanging", "hangingChars"):
        indent.attrib.pop(qn(f"w:{name}"), None)
    indent.set(qn("w:firstLine"), "480")
    indent.set(qn("w:firstLineChars"), "200")


def _remove_outline_level(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    outline_level = p_pr.find(qn("w:outlineLvl"))
    if outline_level is not None:
        p_pr.remove(outline_level)


def _remove_named_bookmarks(doc, names):
    """清理旧版固定目录使用的摘要书签。"""
    matching_ids = set()
    for start in list(doc.element.body.iter(qn("w:bookmarkStart"))):
        if start.get(qn("w:name")) in names:
            matching_ids.add(start.get(qn("w:id")))
            start.getparent().remove(start)
    for end in list(doc.element.body.iter(qn("w:bookmarkEnd"))):
        if end.get(qn("w:id")) in matching_ids:
            end.getparent().remove(end)


def _add_tc_field(paragraph, entry_text):
    """添加隐藏 TC 域，使摘要标题由目录域自动采集。"""
    children = list(paragraph._p)
    for index, child in enumerate(children):
        instruction = "".join(
            element.text or "" for element in child.iter(qn("w:instrText"))
        )
        if not re.search(r"\bTC\s+", instruction):
            continue

        start = index
        while start >= 0:
            fld_char = children[start].find(qn("w:fldChar"))
            if fld_char is not None and fld_char.get(qn("w:fldCharType")) == "begin":
                break
            start -= 1
        end = index
        while end < len(children):
            fld_char = children[end].find(qn("w:fldChar"))
            if fld_char is not None and fld_char.get(qn("w:fldCharType")) == "end":
                break
            end += 1
        for element in children[max(start, 0) : min(end + 1, len(children))]:
            if element.getparent() is paragraph._p:
                paragraph._p.remove(element)
        break

    def hidden_run():
        run = OxmlElement("w:r")
        run_properties = OxmlElement("w:rPr")
        run_properties.append(OxmlElement("w:vanish"))
        run.append(run_properties)
        return run

    begin_run = hidden_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run.append(begin)
    paragraph._p.append(begin_run)

    instruction_run = hidden_run()
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f' TC "{entry_text}" \\f C \\l "1" '
    instruction_run.append(instruction)
    paragraph._p.append(instruction_run)

    end_run = hidden_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)
    paragraph._p.append(end_run)


def _normalize_abstract_navigation(doc):
    chinese_title = _find_paragraph(doc, "内容摘要")
    english_title = _find_paragraph(doc, "ABSTRACT")
    # 用户要求中英文摘要分别成页，英文摘要从中文关键词之后另起一页。
    english_title.paragraph_format.page_break_before = True
    _remove_outline_level(chinese_title)
    _remove_outline_level(english_title)
    _remove_named_bookmarks(doc, {"_AbstractCN", "_AbstractEN"})
    _add_tc_field(chinese_title, "内容摘要")
    _add_tc_field(english_title, "ABSTRACT")

    _set_two_character_first_line_indent(
        _find_paragraph(doc, "关键词：{{KEYWORDS_CN}}")
    )
    _set_two_character_first_line_indent(
        _find_paragraph(doc, "KEYWORDS: {{KEYWORDS_EN}}")
    )


def _remove_reference(sect_pr, tag_name):
    for element in list(sect_pr.findall(qn(f"w:{tag_name}"))):
        sect_pr.remove(element)


def _set_page_numbering(section, fmt=None, start=None):
    sect_pr = section._sectPr
    current = sect_pr.find(qn("w:pgNumType"))
    if current is not None:
        sect_pr.remove(current)
    if fmt is None and start is None:
        return
    element = OxmlElement("w:pgNumType")
    if fmt is not None:
        element.set(qn("w:fmt"), fmt)
    if start is not None:
        element.set(qn("w:start"), str(start))
    sect_pr.append(element)


def _set_run_font(run, name, size_pt):
    run.font.name = name
    run.font.size = Pt(size_pt)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), name)


def _format_paragraph_runs(
    paragraph,
    size_pt,
    east_asia="宋体",
    latin="Times New Roman",
    bold=False,
    alignment=None,
):
    if alignment is not None:
        paragraph.alignment = alignment
    for run in paragraph.runs:
        if not run.text:
            continue
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        r_pr = run._element.get_or_add_rPr()
        r_fonts = r_pr.find(qn("w:rFonts"))
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.insert(0, r_fonts)
        r_fonts.set(qn("w:eastAsia"), east_asia)
        r_fonts.set(qn("w:ascii"), latin)
        r_fonts.set(qn("w:hAnsi"), latin)
        r_fonts.set(qn("w:cs"), latin)


def _normalize_template_typography(doc):
    paragraph_formats = {
        "{{PROJECT_SOURCE}}": (COVER_PROJECT_SOURCE_FONT_SIZE_PT, COVER_HEITI_FONT, COVER_HEITI_FONT, False, WD_ALIGN_PARAGRAPH.CENTER),
        "{{TITLE_CN}}": (COVER_TITLE_CN_FONT_SIZE_PT, COVER_HEITI_FONT, COVER_HEITI_FONT, False, WD_ALIGN_PARAGRAPH.CENTER),
        "{{TITLE_EN}}": (18, "Times New Roman", "Times New Roman", False, WD_ALIGN_PARAGRAPH.CENTER),
        "{{PROJECT_NAME}}": (COVER_SERIES_FONT_SIZE_PT, COVER_HEITI_FONT, COVER_HEITI_FONT, False, WD_ALIGN_PARAGRAPH.CENTER),
        "{{SERIES_NO}}": (COVER_SERIES_FONT_SIZE_PT, COVER_HEITI_FONT, COVER_HEITI_FONT, False, WD_ALIGN_PARAGRAPH.CENTER),
        "{{COVER_BOTTOM_BLOCK}}": (COVER_BOTTOM_FONT_SIZE_PT, COVER_HEITI_FONT, COVER_HEITI_FONT, False, WD_ALIGN_PARAGRAPH.CENTER),
        REPORT_WRITING_UNIT: (COVER_BOTTOM_FONT_SIZE_PT, COVER_HEITI_FONT, COVER_HEITI_FONT, False, WD_ALIGN_PARAGRAPH.CENTER),
        "{{COOPERATOR}}": (COVER_BOTTOM_FONT_SIZE_PT, COVER_HEITI_FONT, COVER_HEITI_FONT, False, WD_ALIGN_PARAGRAPH.CENTER),
        "{{REPORT_DATE}}": (COVER_BOTTOM_FONT_SIZE_PT, COVER_HEITI_FONT, COVER_HEITI_FONT, False, WD_ALIGN_PARAGRAPH.CENTER),
        "内容摘要": (18, "黑体", "Times New Roman", False, WD_ALIGN_PARAGRAPH.CENTER),
        "ABSTRACT": (18, "Times New Roman", "Times New Roman", False, WD_ALIGN_PARAGRAPH.CENTER),
        "目  录": (18, "黑体", "Times New Roman", False, WD_ALIGN_PARAGRAPH.CENTER),
        "关键词：{{KEYWORDS_CN}}": (12, "宋体", "Times New Roman", False, None),
        "KEYWORDS: {{KEYWORDS_EN}}": (12, "Times New Roman", "Times New Roman", False, None),
    }

    approval_start = next(
        i for i, paragraph in enumerate(doc.paragraphs)
        if paragraph.text.strip().startswith("工作单位：")
    )
    approval_end = next(
        i for i, paragraph in enumerate(doc.paragraphs)
        if paragraph.text.strip().startswith("副院长/总工程师")
    )

    metadata_paragraphs = [
        paragraph
        for paragraph in doc.paragraphs
        if re.match(r"^(?:报告编号|分类号|密级)[:：]", paragraph.text.strip())
    ]
    _align_cover_metadata_block(metadata_paragraphs)

    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if re.match(r"^(?:报告编号|分类号|密级)[:：]", text):
            _format_paragraph_runs(
                paragraph,
                COVER_METADATA_FONT_SIZE_PT,
                east_asia=COVER_HEITI_FONT,
                latin=COVER_HEITI_FONT,
                bold=False,
            )
            continue
        if approval_start <= index <= approval_end:
            _format_paragraph_runs(paragraph, 14, bold=True)
            continue
        if text in paragraph_formats:
            _format_paragraph_runs(paragraph, *paragraph_formats[text])


def _add_page_field(paragraph):
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    begin = paragraph.add_run()
    _set_run_font(begin, "宋体", 10.5)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    begin._r.append(fld_begin)

    instruction = paragraph.add_run()
    _set_run_font(instruction, "宋体", 10.5)
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    instruction._r.append(instr_text)

    separate = paragraph.add_run()
    _set_run_font(separate, "宋体", 10.5)
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    separate._r.append(fld_separate)

    cached = paragraph.add_run("i")
    _set_run_font(cached, "宋体", 10.5)

    end = paragraph.add_run()
    _set_run_font(end, "宋体", 10.5)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    end._r.append(fld_end)


def _set_header(paragraph):
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("{{TITLE_CN}}")
    _set_run_font(run, "宋体", 10.5)

    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "auto")


def _merge_abstract_pages(doc):
    chinese_keywords = _find_paragraph(doc, "关键词：{{KEYWORDS_CN}}")
    abstract_title = _find_paragraph(doc, "ABSTRACT")
    current = chinese_keywords._element.getnext()
    while current is not None and current is not abstract_title._element:
        next_element = current.getnext()
        p_pr = current.find(qn("w:pPr"))
        sect_pr = p_pr.find(qn("w:sectPr")) if p_pr is not None else None
        if sect_pr is not None:
            current.getparent().remove(current)
        current = next_element

    # 删除摘要页中的装饰性空行，为中英文摘要保留实际排版空间。
    labels = {
        "内容摘要",
        "{{ABSTRACT_CN}}",
        "关键词：{{KEYWORDS_CN}}",
        "ABSTRACT",
        "{{ABSTRACT_EN}}",
        "KEYWORDS: {{KEYWORDS_EN}}",
    }
    start = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "内容摘要")
    end = next(
        i
        for i, p in enumerate(doc.paragraphs)
        if p.text.strip() == "KEYWORDS: {{KEYWORDS_EN}}"
    )
    for paragraph in list(doc.paragraphs[start : end + 1]):
        if not paragraph.text.strip() and paragraph.text.strip() not in labels:
            _remove_paragraph(paragraph)


def _replace_back_matter_placeholder(doc):
    if any(p.text.strip() == "{{BACK_MATTER}}" for p in doc.paragraphs):
        return
    heading = _find_paragraph(doc, "参考文献")
    placeholder = _find_paragraph(doc, "{{REFERENCES}}")
    _set_text_preserving_first_run(heading, "{{BACK_MATTER}}")
    _remove_paragraph(placeholder)


def _compact_approval_page(doc):
    """删除审批页不承载内容的空行，避免审批签名区溢出到第二页。"""
    start = next(
        i
        for i, paragraph in enumerate(doc.paragraphs)
        if paragraph.text.strip().startswith("工作单位：")
    )
    end = next(
        i
        for i, paragraph in enumerate(doc.paragraphs)
        if paragraph.text.strip().startswith("副院长/总工程师")
    )
    for paragraph in list(doc.paragraphs[start : end + 1]):
        p_pr = paragraph._p.pPr
        has_section = p_pr is not None and p_pr.find(qn("w:sectPr")) is not None
        if not paragraph.text.strip() and not has_section:
            _remove_paragraph(paragraph)


def _add_series_placeholder(doc):
    if any(p.text.strip() == "{{SERIES_NO}}" for p in doc.paragraphs):
        return
    project_name = _find_paragraph(doc, "{{PROJECT_NAME}}")
    next_element = project_name._element.getnext()
    if next_element is None:
        raise ValueError("项目名称后没有可用于系列序号的段落")
    from docx.text.paragraph import Paragraph

    paragraph = Paragraph(next_element, project_name._parent)
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("{{SERIES_NO}}")
    _set_run_font(run, "黑体", 15)


def normalize_template(input_path, output_path):
    doc = Document(input_path)

    if not any(p.text.strip() == "{{PROJECT_SOURCE}}" for p in doc.paragraphs):
        project_source = _find_paragraph(doc, "国家电网有限公司总部科技项目")
        _set_text_preserving_first_run(project_source, "{{PROJECT_SOURCE}}")
    _add_series_placeholder(doc)
    _normalize_cover_bottom_block(doc)
    _compact_approval_page(doc)
    _merge_abstract_pages(doc)
    _normalize_abstract_navigation(doc)
    _replace_back_matter_placeholder(doc)
    _normalize_template_typography(doc)

    # 保存并重新加载，使 python-docx 重新建立合并后的分节列表。
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as temp_file:
        temp_path = temp_file.name
    try:
        doc.save(temp_path)
        doc = Document(temp_path)
    finally:
        os.remove(temp_path)
    if len(doc.sections) != 4:
        raise ValueError(f"规范化后应有 4 个模板分节，实际为 {len(doc.sections)}")

    for section in doc.sections:
        section.page_width = Mm(PAGE_WIDTH_MM)
        section.page_height = Mm(PAGE_HEIGHT_MM)
        section.top_margin = Mm(PAGE_MARGIN_TOP_MM)
        section.bottom_margin = Mm(PAGE_MARGIN_BOTTOM_MM)
        section.left_margin = Mm(PAGE_MARGIN_LEFT_MM)
        section.right_margin = Mm(PAGE_MARGIN_RIGHT_MM)
        section.header_distance = Mm(HEADER_DISTANCE_MM)
        section.footer_distance = Mm(FOOTER_DISTANCE_MM)

    cover, approval, abstract, toc_and_body = doc.sections

    # 封面和审批页不设页眉页脚。
    for section in (cover, approval):
        _remove_reference(section._sectPr, "headerReference")
        _remove_reference(section._sectPr, "footerReference")

    # 摘要页使用罗马数字页码，从 i 开始，不设页眉。
    _remove_reference(abstract._sectPr, "headerReference")
    abstract.footer.is_linked_to_previous = False
    _add_page_field(abstract.footer.paragraphs[0])
    _set_page_numbering(abstract, fmt="lowerRoman", start=1)

    # 目录继承摘要页脚；生成器会在目录和正文间插入分节符。
    toc_and_body.footer.is_linked_to_previous = True
    _set_page_numbering(toc_and_body)

    # 最终文档级分节作为正文分节，显示中文报告名称页眉。
    toc_and_body.header.is_linked_to_previous = False
    _set_header(toc_and_body.header.paragraphs[0])

    doc.save(output_path)


def main():
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    default_path = os.path.join(root, "assets", "report.docx")
    input_path = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else default_path
    output_path = os.path.abspath(sys.argv[2]) if len(sys.argv) > 2 else input_path

    normalize_template(input_path, output_path)
    print(f"模板已规范化: {output_path}")


if __name__ == "__main__":
    main()
