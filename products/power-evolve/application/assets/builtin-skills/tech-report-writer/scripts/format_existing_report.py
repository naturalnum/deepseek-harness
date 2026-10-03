#!/usr/bin/env python3
"""按中国电科院技术报告规范格式化现有 DOCX，并保留全部内容对象。"""

import argparse
import base64
from copy import deepcopy
import os
import re
import subprocess
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.text.paragraph import Paragraph

from config import (
    COVER_HEITI_FONT,
    COVER_METADATA_FONT_SIZE_PT,
    COVER_PROJECT_SOURCE_FONT_SIZE_PT,
    COVER_SERIES_FONT_SIZE_PT,
    COVER_TITLE_CN_FONT_SIZE_PT,
    FOOTER_DISTANCE_MM,
    HEADER_DISTANCE_MM,
    PAGE_HEIGHT_MM,
    PAGE_MARGIN_BOTTOM_MM,
    PAGE_MARGIN_LEFT_MM,
    PAGE_MARGIN_RIGHT_MM,
    PAGE_MARGIN_TOP_MM,
    PAGE_WIDTH_MM,
    REPORT_WRITING_UNIT,
)
from toc_styles import (
    configure_toc_styles,
    ensure_abstract_catalog_marks,
    ensure_figure_table_catalogs,
    mark_figure_table_captions,
)
from normalize_report_template import (
    _add_page_field,
    _add_tc_field,
    _align_cover_metadata_block,
    _remove_outline_level,
    _set_page_numbering,
    _set_text_preserving_first_run,
)
from cover_bottom_layout import (
    extract_and_remove_cover_bottom,
    set_cover_bottom_textbox,
)
from utils import format_cooperator_names, normalize_cooperator_names
from word_field_refresh import refresh_word_fields


BODY_STYLE_NAMES = {
    "Normal",
    "Body Text",
    "Normal (Web)",
    "样式4",
    "本文正文",
    "002二级标题",
    "图",
}
HEADING_2_NAMES = {"Heading 2", "002二级标题"}
HEADING_3_NAMES = {"Heading 3", "003三级标题"}
CAPTION_RE = re.compile(
    r"^(?:图|表)\s*[A-Z0-9]+-\d+[ \t\u3000]+\S"
)
CATALOG_EXPLANATION_LEAD_RE = re.compile(
    r"^(?P<kind>图|表)\s*[A-Z0-9]+-\d+\s*[：:]?\s*"
    r"(?P<verb>应表达|应展示|应给出|给出|展示|显示|表明|说明|反映|可见|"
    r"所示|用于表达|用于展示|用于说明|用于给出)"
)
APPENDIX_TITLE_RE = re.compile(r"^附录\s+[A-Z]\s+\S")
APPENDIX_LEVEL_1_RE = re.compile(r"^[A-Z]\d+\s+\S")
APPENDIX_LEVEL_2_RE = re.compile(r"^[A-Z]\d+\.\d+(?!\.)\s*\S")
APPENDIX_LEVEL_3_RE = re.compile(r"^[A-Z]\d+\.\d+\.\d+(?!\.)\s*\S")
NOTE_RE = re.compile(r"^(?:注(?:\s*\d+)?[:：]|[a-z]\s+\S)")
COVER_METADATA_RE = re.compile(r"^(?:报告编号|分类号|密级)[:：]")
FORMULA_MARKER_RE = re.compile(r"^\[\[TECH_REPORT_MATH_B64:([A-Za-z0-9_=-]+)\]\]$")
INLINE_FORMULA_MARKER_RE = re.compile(
    r"\[\[TECH_REPORT_INLINE_MATH_B64:([A-Za-z0-9_=-]+)\]\]"
)
FORMULA_NUMBER_RE = re.compile(r"^（[A-Z0-9]+-\d+）$")
APPROVAL_LINE_SPACING = 1.25

# 正文段内常见标点/单位/数学技术符号。这里只处理普通 Word 文本 run，
# 不处理 OMML 公式、域代码、上下标引用等具有独立语义的对象。
INLINE_SYMBOL_RE = re.compile(
    r"[!-/:-@\[-`{-~]"
    r"|[\u2000-\u206F\u2100-\u214F\u2190-\u22FF\u2300-\u23FF]"
    r"|[\u2460-\u24FF\u25A0-\u25FF\u3000-\u303F\uFF00-\uFF65]"
    r"|[°℃‰±×÷≈≠≤≥∞√∑∏∫∂∇μΩαβγδΔλφθσ]"
)


def _format_approval_person_names(value: str) -> str:
    """仅用于审批页人员字段：两字中文姓名中间加入一个半角空格。

    示例："张三、李四、王小明" -> "张 三、李 四、王小明"。
    三字及以上中文姓名、英文姓名以及其他页面文本不在此函数处理范围内。
    """
    text = str(value or "")
    return re.sub(
        r"(?<![\u4e00-\u9fff])([\u4e00-\u9fff])([\u4e00-\u9fff])(?![\u4e00-\u9fff])",
        r"\1 \2",
        text,
    )


def _find_paragraph(doc, text):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"文档中未找到段落: {text}")


def _iter_content_paragraphs(doc):
    """Yield body and table-cell paragraphs exactly once, including nested tables."""
    seen = set()

    def emit(paragraph):
        # lxml wrapper objects may be recreated while walking cells; keep the actual
        # element object in the set rather than its transient Python id.
        key = paragraph._p
        if key in seen:
            return None
        seen.add(key)
        return paragraph

    def walk_table(table):
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    item = emit(paragraph)
                    if item is not None:
                        yield item
                for nested in cell.tables:
                    yield from walk_table(nested)

    for paragraph in doc.paragraphs:
        item = emit(paragraph)
        if item is not None:
            yield item
    for table in doc.tables:
        yield from walk_table(table)


def _set_rfonts(element, east_asia="宋体", latin="Times New Roman"):
    r_pr = element.get_or_add_rPr()
    fonts = r_pr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, fonts)
    # 模板、HTML 转换器或 Pandoc 有时会留下 theme/hint 字体属性。
    # 即使已经设置 ascii/hAnsi/eastAsia，这些属性仍可能让单个符号回退到
    # Calibri/Cambria/微软雅黑等字体，从而造成同一段内的符号大小或字形突变。
    for attr_name in (
        "asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "hint"
    ):
        fonts.attrib.pop(qn(f"w:{attr_name}"), None)
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:cs"), latin)


def _set_run_size(run, size_pt):
    """Set both normal and complex-script sizes so Word/WPS cannot fall back differently."""
    run.font.size = Pt(size_pt)
    r_pr = run._r.get_or_add_rPr()
    half_points = str(int(round(size_pt * 2)))
    for name in ("sz", "szCs"):
        size = r_pr.find(qn(f"w:{name}"))
        if size is None:
            size = OxmlElement(f"w:{name}")
            r_pr.append(size)
        size.set(qn("w:val"), half_points)


def _format_run(run, size_pt, east_asia="宋体", latin="Times New Roman", bold=None):
    _set_rfonts(run._element, east_asia=east_asia, latin=latin)
    _set_run_size(run, size_pt)
    if bold is not None:
        run.font.bold = bold


def _format_text_runs(paragraph, size_pt, east_asia="宋体", latin="Times New Roman", bold=None):
    for run in paragraph.runs:
        if run.text:
            _format_run(run, size_pt, east_asia=east_asia, latin=latin, bold=bold)


def _run_has_protected_inline_semantics(run):
    """跳过公式、域、图形和上下标等不应按普通符号重排的 run。"""
    element = run._r
    protected_tags = (
        "w:fldChar", "w:instrText", "w:drawing", "w:object", "w:pict",
        "w:footnoteReference", "w:endnoteReference", "w:commentReference",
    )
    if any(element.find(f".//{qn(tag)}") is not None for tag in protected_tags):
        return True
    if element.find(f".//{qn('m:oMath')}") is not None:
        return True
    r_pr = element.rPr
    if r_pr is not None and r_pr.find(qn("w:vertAlign")) is not None:
        # 典型场景是上标参考文献和上下标符号，保持其语义格式。
        return True
    return False


def _normalize_symbol_run(run, size_pt=12, east_asia="宋体", latin="Times New Roman"):
    """统一普通段内符号的字体槽、字号和非语义性装饰属性。"""
    _set_rfonts(run._element, east_asia=east_asia, latin=latin)
    _set_run_size(run, size_pt)
    # 正文符号不应因源 HTML/run 样式单独变成斜体、下划线或粗体。
    # 上下标/公式已在调用前排除。
    run.font.bold = False
    run.font.italic = False
    run.font.underline = False
    r_pr = run._r.get_or_add_rPr()
    for tag in ("w:position", "w:spacing", "w:w", "w:kern"):
        node = r_pr.find(qn(tag))
        if node is not None:
            r_pr.remove(node)


def _is_normal_body_text_paragraph(paragraph):
    text = paragraph.text.strip()
    if not text:
        return False
    style_name = paragraph.style.name if paragraph.style is not None else ""
    if style_name.lower().startswith("toc "):
        return False
    if style_name in {"Heading 1", "Caption"} | HEADING_2_NAMES | HEADING_3_NAMES:
        return False
    if (
        APPENDIX_TITLE_RE.match(text) or APPENDIX_LEVEL_1_RE.match(text)
        or APPENDIX_LEVEL_2_RE.match(text) or APPENDIX_LEVEL_3_RE.match(text)
        or CAPTION_RE.match(text) or NOTE_RE.match(text)
        or text == "参考文献"
    ):
        return False
    if paragraph.alignment in (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT) and len(text) < 120:
        return False
    return True


def _normalize_inline_symbol_formatting(doc):
    """在最终保存前统一正文/表格普通文本中的段内符号格式。

    该步骤故意放在行内公式替换之后执行，修复后插入对象或源文档 run
    留下的 theme/hint 字体和异常字号。公式 OMML、域代码、上/下标引用不处理。
    """
    try:
        body_start = next(
            i for i, paragraph in enumerate(doc.paragraphs)
            if paragraph.style.name == "Heading 1" and re.match(r"^1\s+", paragraph.text.strip())
        )
    except StopIteration:
        body_start = 0

    changed = 0
    for paragraph in doc.paragraphs[body_start:]:
        if not _is_normal_body_text_paragraph(paragraph):
            continue
        for run in paragraph.runs:
            if (
                not run.text
                or INLINE_SYMBOL_RE.search(run.text) is None
                or _run_has_protected_inline_semantics(run)
            ):
                continue
            _normalize_symbol_run(run, 12)
            changed += 1

    # 表格正文统一为五号，符号也不得继承外部模板字体。
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if (
                            not run.text
                            or INLINE_SYMBOL_RE.search(run.text) is None
                            or _run_has_protected_inline_semantics(run)
                        ):
                            continue
                        _normalize_symbol_run(run, 10.5)
                        changed += 1
    return changed


def _set_first_line_chars(paragraph, chars=200, twips=480):
    p_pr = paragraph._p.get_or_add_pPr()
    indent = p_pr.find(qn("w:ind"))
    if indent is None:
        indent = OxmlElement("w:ind")
        p_pr.append(indent)
    for name in ("left", "leftChars", "right", "rightChars", "hanging", "hangingChars"):
        indent.attrib.pop(qn(f"w:{name}"), None)
    indent.set(qn("w:firstLine"), str(twips))
    indent.set(qn("w:firstLineChars"), str(chars))


def _clear_indent(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    indent = p_pr.find(qn("w:ind"))
    if indent is not None:
        p_pr.remove(indent)


def _set_zero_first_line_indent(paragraph):
    """强制首行缩进为 0，避免继承正文样式中的两字符缩进。"""
    _clear_indent(paragraph)
    p_pr = paragraph._p.get_or_add_pPr()
    indent = OxmlElement("w:ind")
    indent.set(qn("w:firstLine"), "0")
    indent.set(qn("w:firstLineChars"), "0")
    p_pr.append(indent)
    paragraph.paragraph_format.first_line_indent = Pt(0)


def _element_plain_text(element):
    """读取块级 XML 中非数学公式的普通 Word 文本。"""
    return "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()


def _element_contains_math(element):
    return (
        element.find(f".//{qn('m:oMath')}") is not None
        or element.find(f".//{qn('m:oMathPara')}") is not None
    )


def _is_display_formula_block(element):
    """识别独立公式块，避免把正文中的行内公式误判为显示公式。"""
    if element.tag == qn("w:tbl"):
        return _element_contains_math(element)
    if element.tag != qn("w:p") or not _element_contains_math(element):
        return False
    plain_text = _element_plain_text(element)
    return not plain_text or FORMULA_NUMBER_RE.fullmatch(plain_text) is not None


def _remove_first_line_indent_after_formulas(doc):
    """将每个独立公式后的首个实际文字段落设为首行不缩进。

    空段、单独的公式编号和连续公式不计为“公式后的文字段落”。
    若公式后直接出现非公式表格，则不跨过该表格继续修改后续段落。
    """
    body = doc.element.body
    children = list(body)
    paragraph_map = {id(paragraph._p): paragraph for paragraph in doc.paragraphs}
    changed = 0

    for index, child in enumerate(children):
        if not _is_display_formula_block(child):
            continue

        for candidate in children[index + 1 :]:
            if candidate.tag == qn("w:tbl"):
                if _is_display_formula_block(candidate):
                    continue
                break
            if candidate.tag != qn("w:p"):
                continue
            if _is_display_formula_block(candidate):
                continue

            text = _element_plain_text(candidate)
            if not text:
                continue
            if FORMULA_NUMBER_RE.fullmatch(text):
                continue

            paragraph = paragraph_map.get(id(candidate))
            if paragraph is not None:
                _set_zero_first_line_indent(paragraph)
                changed += 1
            break

    return changed


def _format_approval_paragraph(paragraph, index, label, value):
    """Format one approval-page row with a shared value column."""
    paragraph.clear()
    if label or value:
        text = label
        if value:
            if index == 1:
                text += value
            else:
                text += f"\t{value.replace(chr(10), chr(10) + chr(9))}"
        paragraph.add_run(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Mm(
        32 if index == 1 else (14 if index >= 7 else 0)
    )
    paragraph.paragraph_format.right_indent = None
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.line_spacing = APPROVAL_LINE_SPACING
    p_pr = paragraph._p.get_or_add_pPr()
    tabs = p_pr.find(qn("w:tabs"))
    if tabs is not None:
        p_pr.remove(tabs)
    if value and index != 1:
        paragraph.paragraph_format.tab_stops.add_tab_stop(Mm(32))
    _clear_indent(paragraph)
    if index == 1:
        paragraph.paragraph_format.left_indent = Mm(32)
    if index >= 7:
        paragraph.paragraph_format.left_indent = Mm(14)
    paragraph.paragraph_format.space_after = Pt(0)
    _format_text_runs(paragraph, 14, bold=True)


def _insert_approval_blank_line(anchor):
    """Insert a visible blank approval-page line immediately after *anchor*."""
    element = OxmlElement("w:p")
    anchor._p.addnext(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.add_run(" ")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Mm(0)
    paragraph.paragraph_format.right_indent = None
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = APPROVAL_LINE_SPACING
    _format_text_runs(paragraph, 14, bold=True)
    return paragraph


def _remove_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)


def _has_embedded_object(paragraph):
    return any(
        paragraph._p.findall(f".//{qn(tag)}")
        for tag in ("w:drawing", "w:object", "w:pict")
    )


def _has_numbering(paragraph):
    p_pr = paragraph._p.pPr
    return p_pr is not None and p_pr.find(qn("w:numPr")) is not None


def _is_english_dominant(text):
    letters = len(re.findall(r"[A-Za-z]", text))
    chinese = len(re.findall(r"[\u4e00-\u9fff]", text))
    return letters > chinese * 2 and letters >= 8


def _cover_boundaries(doc):
    paragraphs = doc.paragraphs
    approval_start = next((i for i, p in enumerate(paragraphs)
                           if p.text.strip().startswith("工作单位：")), None)
    if approval_start is None:
        raise ValueError("封面结构缺少审批页“工作单位：”标识")
    cover_section_end_index = next((i for i, p in enumerate(paragraphs[:approval_start])
        if p._p.pPr is not None and p._p.pPr.find(qn("w:sectPr")) is not None), None)
    if cover_section_end_index is None:
        raise ValueError("封面结构缺少封面与审批页之间的分节符")
    source_index = next((i for i, p in enumerate(paragraphs[:cover_section_end_index])
        if any(b.get(qn("w:name")) == "PEReportSource"
               for b in p._p.iter(qn("w:bookmarkStart")))), None)
    if source_index is None:
        # Older reports have no semantic bookmark. Locate their source through
        # the three cover metadata rows, never through the source's wording.
        metadata = [i for i, p in enumerate(paragraphs[:cover_section_end_index])
                    if COVER_METADATA_RE.match(p.text.strip())]
        if len(metadata) != 3:
            raise ValueError("封面结构缺少报告编号、分类号或密级")
        source_index = next((i for i in range(metadata[-1] + 1, cover_section_end_index)
                             if paragraphs[i].text.strip()), None)
    if source_index is None or source_index + 1 >= cover_section_end_index:
        raise ValueError("封面结构缺少项目来源或报告标题")
    cover_content_end_index = next(
        (
            i
            for i, paragraph in enumerate(
                paragraphs[source_index + 1 : cover_section_end_index],
                source_index + 1,
            )
            if paragraph.text.strip().startswith(REPORT_WRITING_UNIT)
        ),
        cover_section_end_index,
    )
    series_index = next(
        (
            i
            for i, paragraph in enumerate(
                paragraphs[source_index + 1 : cover_content_end_index],
                source_index + 1,
            )
            if paragraph.text.strip().startswith("《")
        ),
        cover_content_end_index,
    )
    return source_index, series_index, cover_content_end_index, approval_start


def _extract_report_title(doc):
    source_index, _, cover_end_index, _ = _cover_boundaries(doc)
    title_lines = []
    for paragraph in doc.paragraphs[source_index + 1 : cover_end_index]:
        text = paragraph.text.strip()
        if not text:
            continue
        if _is_english_dominant(text):
            break
        title_lines.append(text)
    if not title_lines:
        raise ValueError("封面中未找到技术报告中文名称")
    return "".join(title_lines)


def _format_heading(paragraph, level):
    sizes = {1: 15, 2: 14, 3: 12}
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = paragraph.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.5
    pf.first_line_indent = None
    pf.left_indent = None
    pf.right_indent = None
    pf.keep_with_next = True
    pf.keep_together = True
    pf.page_break_before = level == 1
    _clear_indent(paragraph)
    _format_text_runs(paragraph, sizes[level], east_asia="黑体", bold=False)


def _format_body(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = paragraph.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.keep_with_next = False
    pf.keep_together = False
    _disable_grid_snap(paragraph)
    _set_first_line_chars(paragraph)
    _format_text_runs(paragraph, 12, bold=False)


def _format_reference(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = paragraph.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    _disable_grid_snap(paragraph)
    _set_first_line_chars(paragraph)
    _format_text_runs(paragraph, 12, bold=False)


def _disable_grid_snap(paragraph):
    # A legacy section grid rounds CJK fallback-font lines beyond the specified
    # 1.5 spacing. Keep the template grid for other content and override this text.
    properties = paragraph._p.get_or_add_pPr()
    snap = properties.find(qn("w:snapToGrid"))
    if snap is None:
        snap = OxmlElement("w:snapToGrid")
        properties.append(snap)
    snap.set(qn("w:val"), "0")


def _format_special_title(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = paragraph.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.5
    pf.first_line_indent = None
    pf.page_break_before = True
    pf.keep_with_next = True
    _clear_indent(paragraph)
    _format_text_runs(paragraph, 15, east_asia="黑体", bold=False)


def _format_note(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = paragraph.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = None
    _clear_indent(paragraph)
    _format_text_runs(paragraph, 10.5, bold=False)


def _format_caption(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = paragraph.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = None
    pf.keep_together = True
    # Table captions precede their tables and must not be stranded at the
    # bottom of a page. Figure captions follow figures, so keep-with-next
    # would incorrectly bind them to the following body paragraph.
    pf.keep_with_next = paragraph.text.strip().startswith("表")
    _clear_indent(paragraph)
    _format_text_runs(paragraph, 10.5, east_asia="黑体", bold=False)


def _replace_visible_prefix(paragraph, matched_length: int, replacement: str) -> None:
    """Replace a visible paragraph prefix without flattening the rest of OOXML."""
    text_nodes = list(paragraph._p.iter(qn("w:t")))
    remaining = matched_length
    inserted = False
    for node in text_nodes:
        value = node.text or ""
        if remaining <= 0:
            break
        take = min(len(value), remaining)
        suffix = value[take:]
        if not inserted:
            node.text = replacement + suffix
            inserted = True
        else:
            node.text = suffix
        remaining -= take


def _normalize_catalog_explanation_paragraphs(doc) -> int:
    """Rewrite number-led figure/table explanatory prose to non-caption wording."""
    changed = 0
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        match = CATALOG_EXPLANATION_LEAD_RE.match(text)
        if match is None:
            continue
        kind = match.group("kind")
        verb = match.group("verb")
        if verb == "所示":
            replacement = f"如该{kind}所示"
        elif verb in {"显示", "表明", "反映", "可见"}:
            replacement = f"{'上述趋势' if kind == '图' else '上述结果'}{verb}"
        else:
            replacement = f"该{kind}{verb}"
        _replace_visible_prefix(paragraph, match.end(), replacement)
        changed += 1
    return changed


def _format_style(style, size_pt, east_asia, latin="Times New Roman", bold=False):
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    r_pr = style.element.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = OxmlElement("w:rPr")
        style.element.append(r_pr)
    fonts = r_pr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:cs"), latin)


def _remove_abstract_section_break(doc):
    chinese_keywords = _find_paragraph(doc, next(
        p.text.strip() for p in doc.paragraphs if p.text.strip().startswith("关键词：")
    ))
    english_title = _find_paragraph(doc, "ABSTRACT")
    current = chinese_keywords._element.getnext()
    while current is not None and current is not english_title._element:
        next_element = current.getnext()
        p_pr = current.find(qn("w:pPr"))
        sect_pr = p_pr.find(qn("w:sectPr")) if p_pr is not None else None
        if sect_pr is not None:
            p_pr.remove(sect_pr)
            if not "".join(text.text or "" for text in current.iter(qn("w:t"))).strip():
                current.getparent().remove(current)
            return True
        current = next_element
    return False


def _remove_section_break_before(doc, target_text):
    target = next((p for p in doc.paragraphs if p.text.strip() == target_text), None)
    if target is None:
        return False
    current = target._element.getprevious()
    while current is not None:
        p_pr = current.find(qn("w:pPr"))
        sect_pr = p_pr.find(qn("w:sectPr")) if p_pr is not None else None
        if sect_pr is not None:
            p_pr.remove(sect_pr)
            if not "".join(text.text or "" for text in current.iter(qn("w:t"))).strip():
                current.getparent().remove(current)
            return True
        if "".join(text.text or "" for text in current.iter(qn("w:t"))).strip():
            break
        current = current.getprevious()
    return False


def _clear_part(part):
    for paragraph in part.paragraphs:
        paragraph.clear()


def _set_header(paragraph, title):
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    _format_run(run, 10.5)
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "auto")


def _configure_sections(doc, title):
    if len(doc.sections) != 5:
        raise ValueError(f"规范化后应有 5 个分节，实际为 {len(doc.sections)}")

    for section in doc.sections:
        section.page_width = Mm(PAGE_WIDTH_MM)
        section.page_height = Mm(PAGE_HEIGHT_MM)
        section.top_margin = Mm(PAGE_MARGIN_TOP_MM)
        section.bottom_margin = Mm(PAGE_MARGIN_BOTTOM_MM)
        section.left_margin = Mm(PAGE_MARGIN_LEFT_MM)
        section.right_margin = Mm(PAGE_MARGIN_RIGHT_MM)
        section.header_distance = Mm(HEADER_DISTANCE_MM)
        section.footer_distance = Mm(FOOTER_DISTANCE_MM)
        section.start_type = WD_SECTION_START.NEW_PAGE

    cover, approval, abstract, toc, body = doc.sections

    for section in (cover, approval, abstract, toc):
        section.header.is_linked_to_previous = False
        _clear_part(section.header)

    for section in (cover, approval):
        section.footer.is_linked_to_previous = False
        _clear_part(section.footer)
        _set_page_numbering(section)

    abstract.footer.is_linked_to_previous = False
    _add_page_field(abstract.footer.paragraphs[0])
    _set_page_numbering(abstract, fmt="lowerRoman", start=1)

    toc.footer.is_linked_to_previous = True
    _set_page_numbering(toc, fmt="lowerRoman")

    body.header.is_linked_to_previous = False
    _set_header(body.header.paragraphs[0], title)
    body.footer.is_linked_to_previous = False
    _add_page_field(body.footer.paragraphs[0])
    _set_page_numbering(body, fmt="decimal", start=1)

def _configure_front_matter(doc):
    bottom_lines = extract_and_remove_cover_bottom(doc)
    if len(bottom_lines) < 2:
        raise ValueError("封面页底缺少报告发布日期")

    report_date = bottom_lines[-1]
    cooperator_lines = normalize_cooperator_names(bottom_lines[1:-1])
    set_cover_bottom_textbox(
        doc,
        [REPORT_WRITING_UNIT, *cooperator_lines, report_date],
    )

    source_index, series_index, cover_end_index, approval_start = _cover_boundaries(doc)

    metadata_paragraphs = [
        paragraph
        for paragraph in doc.paragraphs[:source_index]
        if COVER_METADATA_RE.match(paragraph.text.strip())
    ]
    for paragraph in metadata_paragraphs:
        _format_text_runs(
            paragraph,
            COVER_METADATA_FONT_SIZE_PT,
            east_asia=COVER_HEITI_FONT,
            latin=COVER_HEITI_FONT,
            bold=False,
        )
    _align_cover_metadata_block(metadata_paragraphs)

    source = doc.paragraphs[source_index]
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _format_text_runs(
        source,
        COVER_PROJECT_SOURCE_FONT_SIZE_PT,
        east_asia=COVER_HEITI_FONT,
        latin=COVER_HEITI_FONT,
        bold=False,
    )

    title_area = doc.paragraphs[source_index + 1 : series_index]
    english_title_offset = next(
        (
            i
            for i, paragraph in enumerate(title_area)
            if _is_english_dominant(paragraph.text.strip())
        ),
        None,
    )
    title_end = len(title_area) if english_title_offset is None else english_title_offset + 1

    for paragraph in title_area[:title_end]:
        text = paragraph.text.strip()
        if not text:
            continue
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if _is_english_dominant(text):
            _format_text_runs(
                paragraph,
                16,
                east_asia="Times New Roman",
                latin="Times New Roman",
                bold=False,
            )
        else:
            _format_text_runs(
                paragraph,
                COVER_TITLE_CN_FONT_SIZE_PT,
                east_asia=COVER_HEITI_FONT,
                latin=COVER_HEITI_FONT,
                bold=False,
            )

    # Long report and project names can otherwise push the institution/date
    # block onto a second cover page. Preserve the template's visual gap while
    # compressing its empty spacer paragraphs.
    for paragraph in title_area:
        if paragraph.text.strip():
            continue
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = Pt(3)

    project_and_series = list(title_area[title_end:])
    project_and_series.extend(doc.paragraphs[series_index:cover_end_index])
    for paragraph in project_and_series:
        if not paragraph.text.strip():
            continue
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _format_text_runs(
            paragraph,
            COVER_SERIES_FONT_SIZE_PT,
            east_asia=COVER_HEITI_FONT,
            latin=COVER_HEITI_FONT,
            bold=False,
        )

    approval_section_end = next(
        i
        for i, paragraph in enumerate(doc.paragraphs[approval_start:], approval_start)
        if (
            paragraph._p.pPr is not None
            and paragraph._p.pPr.find(qn("w:sectPr")) is not None
        )
    )
    source_approval_paragraphs = doc.paragraphs[
        approval_start:approval_section_end
    ]

    def approval_value(label):
        pattern = re.compile(rf"^{re.escape(label)}\s*[:：]?\s*")
        for paragraph in source_approval_paragraphs:
            text = paragraph.text.strip()
            if pattern.match(text):
                return pattern.sub("", text, count=1).lstrip("\t ")
        return ""

    approval_values = {
        "work_unit": REPORT_WRITING_UNIT,
        "cooperator": format_cooperator_names(cooperator_lines),
        "work_period": approval_value("工作时间"),
        "leader": _format_approval_person_names(approval_value("项目负责人")),
        "staff": _format_approval_person_names(approval_value("工作人员")),
        "writer": _format_approval_person_names(approval_value("报告编写")),
    }
    approval_texts = [
        ("工作单位：", approval_values["work_unit"]),
        ("", approval_values["cooperator"]),
        ("工作时间：", approval_values["work_period"]),
        ("项目负责人：", approval_values["leader"]),
        ("工作人员：", approval_values["staff"]),
        ("报告编写：", approval_values["writer"]),
        (REPORT_WRITING_UNIT, ""),
        ("报告审核：", ""),
        ("学术技术分委会：", ""),
        ("部　　门：", ""),
        ("科技部/咨询业务部（格式审核）：", ""),
        ("学术技术委员会：", ""),
        ("报告批准", ""),
        ("副院长/总工程师/院长助理/副总工程师：", ""),
    ]

    # Approval-page templates in circulation are not identical: some omit the
    # final approver row, while others contain extra spacer paragraphs. Rebuild
    # the normalized rows before the existing section-break paragraph instead
    # of depending on a particular final label being present.
    approval_boundary = doc.paragraphs[approval_section_end]
    for paragraph in source_approval_paragraphs:
        _remove_paragraph(paragraph)
    approval_paragraphs = []
    for _ in approval_texts:
        element = OxmlElement("w:p")
        approval_boundary._p.addprevious(element)
        approval_paragraphs.append(Paragraph(element, approval_boundary._parent))

    for index, (paragraph, (label, value)) in enumerate(
        zip(approval_paragraphs, approval_texts)
    ):
        # 合作单位没有字段标签，应以工作单位的取值起点对齐，而不是
        # 从版心左侧起排。其余审批信息仍由制表位定位取值。
        _format_approval_paragraph(paragraph, index, label, value)
    _insert_approval_blank_line(approval_paragraphs[5])
    if approval_values["cooperator"]:
        _insert_approval_blank_line(approval_paragraphs[1])

    chinese_title = _find_paragraph(doc, "内容摘要")
    chinese_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _format_text_runs(chinese_title, 18, east_asia="黑体", bold=False)
    _remove_outline_level(chinese_title)
    _add_tc_field(chinese_title, "内容摘要")

    english_title = _find_paragraph(doc, "ABSTRACT")
    english_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _format_text_runs(
        english_title,
        18,
        east_asia="Times New Roman",
        latin="Times New Roman",
        bold=False,
    )
    # 中文摘要及关键词结束后，英文摘要另起一页。
    english_title.paragraph_format.page_break_before = True
    _remove_outline_level(english_title)
    _add_tc_field(english_title, "ABSTRACT")

    chinese_keywords = next(p for p in doc.paragraphs if p.text.strip().startswith("关键词："))
    english_keywords = next(p for p in doc.paragraphs if p.text.strip().startswith("KEYWORDS:"))
    chinese_title_index = next(
        i for i, paragraph in enumerate(doc.paragraphs) if paragraph._p is chinese_title._p
    )
    chinese_keywords_index = next(
        i for i, paragraph in enumerate(doc.paragraphs) if paragraph._p is chinese_keywords._p
    )
    english_title_index = next(
        i for i, paragraph in enumerate(doc.paragraphs) if paragraph._p is english_title._p
    )
    english_keywords_index = next(
        i for i, paragraph in enumerate(doc.paragraphs) if paragraph._p is english_keywords._p
    )

    abstract_ranges = (
        (chinese_title_index + 1, chinese_keywords_index + 1, "宋体", "Times New Roman"),
        (
            english_title_index + 1,
            english_keywords_index + 1,
            "Times New Roman",
            "Times New Roman",
        ),
    )
    for start, end, east_asia, latin in abstract_ranges:
        for paragraph in doc.paragraphs[start:end]:
            if not paragraph.text.strip():
                continue
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            _set_first_line_chars(paragraph)
            _format_text_runs(
                paragraph,
                12,
                east_asia=east_asia,
                latin=latin,
                bold=False,
            )

    toc_title = next(
        p
        for p in doc.paragraphs
        if re.sub(r"\s+", "", p.text.strip()) == "目录"
    )
    if toc_title.runs:
        toc_title.runs[0].text = "目  录"
        for run in toc_title.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        toc_title.add_run("目  录")
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _format_text_runs(toc_title, 18, east_asia="黑体", bold=False)


def _configure_toc(doc):
    configure_toc_styles(doc)
    # Content generation should already avoid number-led explanatory prose, but
    # existing reports and legacy drafts may still contain it. Normalize before
    # catalog marking so only true caption paragraphs can receive F/T TC fields.
    _normalize_catalog_explanation_paragraphs(doc)
    ensure_abstract_catalog_marks(doc)
    mark_figure_table_captions(doc)
    ensure_figure_table_catalogs(doc)
    found_toc = False
    for paragraph in doc.paragraphs:
        instruction = "".join(
            element.text or "" for element in paragraph._p.iter(qn("w:instrText"))
        )
        if not re.search(r"\bTOC\b", instruction):
            continue
        found_toc = True
        for element in paragraph._p.iter(qn("w:instrText")):
            if re.search(r"\bTOC\b", element.text or ""):
                element.text = ' TOC \\o "1-3" \\f C \\h \\z \\u '
        for field in paragraph._p.iter(qn("w:fldChar")):
            if field.get(qn("w:fldCharType")) == "begin":
                field.set(qn("w:dirty"), "true")
                break
        break
    if not found_toc:
        raise ValueError("文档中未找到 TOC 目录域")

    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    for paragraph in doc.paragraphs:
        if paragraph.style.name.lower().startswith("toc "):
            paragraph.paragraph_format.line_spacing = 1.0
            _format_text_runs(paragraph, 12, bold=False)


def _configure_body(doc):
    heading_styles = [
        ("Heading 1", 1),
        ("Heading 2", 2),
        ("Heading 3", 3),
        ("003三级标题", 3),
    ]
    for style_name, level in heading_styles:
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        _format_style(style, {1: 15, 2: 14, 3: 12}[level], "黑体")
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.page_break_before = level == 1

    body_start = next(
        i
        for i, paragraph in enumerate(doc.paragraphs)
        if paragraph.style.name == "Heading 1" and re.match(r"^1\s+", paragraph.text.strip())
    )
    reference_start = next(
        (
            i
            for i, paragraph in enumerate(doc.paragraphs)
            if paragraph.text.strip() == "参考文献"
        ),
        len(doc.paragraphs),
    )

    for index, paragraph in enumerate(doc.paragraphs[body_start:], start=body_start):
        text = paragraph.text.strip()
        if not text:
            continue
        if text == "参考文献":
            _format_special_title(paragraph)
            continue
        if APPENDIX_TITLE_RE.match(text):
            _format_special_title(paragraph)
            continue
        if index > reference_start:
            _format_reference(paragraph)
            continue
        if paragraph.style.name == "Heading 1":
            _format_heading(paragraph, 1)
            continue
        if APPENDIX_LEVEL_1_RE.match(text):
            _format_heading(paragraph, 1)
            continue
        if (
            paragraph.style.name in HEADING_3_NAMES
            or APPENDIX_LEVEL_3_RE.match(text)
        ):
            _format_heading(paragraph, 3)
            continue
        if paragraph.style.name in HEADING_2_NAMES or APPENDIX_LEVEL_2_RE.match(text):
            _format_heading(paragraph, 2)
            continue
        if paragraph.style.name == "Caption" or CAPTION_RE.match(text):
            _format_caption(paragraph)
            continue
        if NOTE_RE.match(text):
            _format_note(paragraph)
            continue
        if paragraph.alignment in (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT) and len(text) < 120:
            _format_text_runs(paragraph, 10.5, bold=False)
            continue
        if paragraph.style.name in BODY_STYLE_NAMES or paragraph.style.name not in {
            "Heading 1",
            "Heading 2",
            "Heading 3",
            "003三级标题",
        }:
            _format_body(paragraph)

    for table in doc.tables:
        for row_index, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:cantSplit")) is None:
                tr_pr.append(OxmlElement("w:cantSplit"))
            if row_index == 0 and tr_pr.find(qn("w:tblHeader")) is None:
                tr_pr.append(OxmlElement("w:tblHeader"))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if _has_embedded_object(paragraph):
                        continue
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.first_line_indent = None
                    _format_text_runs(paragraph, 10.5, bold=False)


def _pandoc_math_elements(formulas):
    if not formulas:
        return []
    markdown = "\n\n".join(f"$$\n{formula}\n$$" for formula in formulas)
    with tempfile.TemporaryDirectory(prefix="tech_report_math_") as temp_dir:
        markdown_path = os.path.join(temp_dir, "formulas.md")
        docx_path = os.path.join(temp_dir, "formulas.docx")
        with open(markdown_path, "w", encoding="utf-8") as file:
            file.write(markdown)
        subprocess.run(
            [
                "pandoc",
                "--from=markdown+tex_math_dollars",
                "--to=docx",
                "--output",
                docx_path,
                markdown_path,
            ],
            check=True,
            capture_output=True,
            text=True,
        )
        converted = Document(docx_path)
        math_elements = [
            deepcopy(element)
            for paragraph in converted.paragraphs
            for element in paragraph._p.iter(qn("m:oMath"))
        ]
    if len(math_elements) != len(formulas):
        raise ValueError(
            f"公式转换数量不一致：源公式 {len(formulas)} 个，Word 公式 {len(math_elements)} 个"
        )
    return math_elements


def _display_formula_rows(formula, row_budget=72):
    """Return safe visual rows using line boundaries already in the source."""
    source_lines = [line.strip() for line in formula.splitlines() if line.strip()]
    if len(source_lines) < 2:
        return [formula]
    if r"\begin{cases}" in formula:
        return [formula]
    if any(line.startswith("{") for line in source_lines):
        return [formula]
    if any(
        line.count(r"\left") != line.count(r"\right")
        for line in source_lines
    ):
        return [formula]

    should_wrap = len(source_lines) >= 3 or (
        len(source_lines) == 2
        and (
            source_lines[0].endswith(r"\quad")
            or source_lines[1].startswith(r"-\mu")
        )
    )
    if not should_wrap:
        return [formula]

    atomic_lines = []
    for source_line in reversed(source_lines):
        starts_with_sum = re.match(r"^[+-]?\\sum", source_line)
        ends_with_sum = re.search(
            r"\\sum(?:_\{[^{}]*\})?(?:\^\{[^{}]*\})?\s*$",
            source_line,
        )
        if (starts_with_sum or ends_with_sum) and atomic_lines:
            atomic_lines[0] = f"{source_line} {atomic_lines[0]}"
        else:
            atomic_lines.insert(0, source_line)

    if len(atomic_lines) == 2:
        return atomic_lines

    rows = []
    current = ""
    for source_line in atomic_lines:
        candidate = f"{current} {source_line}".strip()
        if current and len(candidate) > row_budget:
            rows.append(current)
            current = source_line
        else:
            current = candidate
    if current:
        rows.append(current)

    if len(rows) < 2:
        return [formula]
    return rows


def _display_formula_size(formula):
    compact_length = len(re.sub(r"\s+", "", formula))
    if compact_length >= 220:
        return 7.5
    if compact_length >= 160:
        return 8.0
    if compact_length >= 125:
        return 8.5
    if compact_length >= 100:
        return 9.0
    return 10.5


def _format_math_element(math_element, size_pt):
    half_points = str(int(round(size_pt * 2)))
    for math_run in math_element.iter(qn("m:r")):
        run_pr = math_run.find(qn("w:rPr"))
        if run_pr is None:
            run_pr = OxmlElement("w:rPr")
            math_properties = math_run.find(qn("m:rPr"))
            if math_properties is None:
                math_run.insert(0, run_pr)
            else:
                math_properties.addnext(run_pr)
        fonts = run_pr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            run_pr.insert(0, fonts)
        for name in ("ascii", "hAnsi", "cs"):
            fonts.set(qn(f"w:{name}"), "Cambria Math")
        for name in ("sz", "szCs"):
            size = run_pr.find(qn(f"w:{name}"))
            if size is None:
                size = OxmlElement(f"w:{name}")
                run_pr.append(size)
            size.set(qn("w:val"), half_points)
    return math_element


def _replace_inline_formula_markers(doc):
    paragraphs = []
    formulas = []
    for paragraph in _iter_content_paragraphs(doc):
        matches = list(INLINE_FORMULA_MARKER_RE.finditer(paragraph.text))
        if not matches:
            continue
        paragraphs.append((paragraph, matches))
        formulas.extend(
            base64.urlsafe_b64decode(match.group(1).encode("ascii")).decode("utf-8")
            for match in matches
        )

    math_elements = iter(_pandoc_math_elements(formulas))
    for paragraph, matches in paragraphs:
        original_text = paragraph.text
        original_r_pr = next(
            (
                deepcopy(run._r.rPr)
                for run in paragraph.runs
                if run._r.rPr is not None
            ),
            None,
        )
        paragraph.clear()
        cursor = 0
        for match in matches:
            if match.start() > cursor:
                run = paragraph.add_run(original_text[cursor : match.start()])
                if original_r_pr is not None:
                    run._r.insert(0, deepcopy(original_r_pr))
            paragraph._p.append(
                _format_math_element(next(math_elements), 12)
            )
            cursor = match.end()
        if cursor < len(original_text):
            run = paragraph.add_run(original_text[cursor:])
            if original_r_pr is not None:
                run._r.insert(0, deepcopy(original_r_pr))


def _set_cell_width(cell, width_mm):
    cell.width = Mm(width_mm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_width = tc_pr.find(qn("w:tcW"))
    if tc_width is None:
        tc_width = OxmlElement("w:tcW")
        tc_pr.append(tc_width)
    tc_width.set(qn("w:w"), str(int(Mm(width_mm).twips)))
    tc_width.set(qn("w:type"), "dxa")


def _set_formula_table_widths(table):
    """同步编号显示公式的固定列网格、表格总宽度与单元格宽度。"""
    table.autofit = False
    column_widths = (0, 146, 14)
    for column, width in zip(table.columns, column_widths):
        column.width = Mm(width)
    table_width = table._tbl.tblPr.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table._tbl.tblPr.append(table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(sum(int(Mm(width).twips) for width in column_widths)))
    for cell, width in zip(table.rows[0].cells, column_widths):
        _set_cell_width(cell, width)


def _repair_existing_formula_tables(doc):
    """校准单行三列编号 OMML 表格，保留含文字标签或其他对象的数据表。"""
    changed = 0
    for table in doc.tables:
        if len(table.rows) != 1 or len(table.columns) != 3:
            continue
        left, equation, number = table.rows[0].cells
        if _element_plain_text(left._tc) or _element_plain_text(equation._tc):
            continue
        if not _element_contains_math(equation._tc) or not FORMULA_NUMBER_RE.fullmatch(number.text.strip()):
            continue
        if any(_has_embedded_object(paragraph) for cell in (left, number) for paragraph in cell.paragraphs):
            continue
        _set_formula_table_widths(table)
        changed += 1
    return changed


def _remove_table_borders(table):
    table_pr = table._tbl.tblPr
    borders = table_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_pr.append(borders)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{name}"))
        if edge is None:
            edge = OxmlElement(f"w:{name}")
            borders.append(edge)
        edge.set(qn("w:val"), "nil")


def _set_cell_margins_zero(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name in ("top", "left", "bottom", "right"):
        margin = margins.find(qn(f"w:{name}"))
        if margin is None:
            margin = OxmlElement(f"w:{name}")
            margins.append(margin)
        margin.set(qn("w:w"), "0")
        margin.set(qn("w:type"), "dxa")


def _replace_formula_markers(doc):
    marker_paragraphs = []
    formulas = []
    for paragraph in _iter_content_paragraphs(doc):
        match = FORMULA_MARKER_RE.fullmatch(paragraph.text.strip())
        if not match:
            continue
        formulas.append(
            base64.urlsafe_b64decode(match.group(1).encode("ascii")).decode("utf-8")
        )
        marker_paragraphs.append(paragraph)

    formula_rows = [_display_formula_rows(formula) for formula in formulas]
    formula_sizes = [_display_formula_size(formula) for formula in formulas]
    math_elements = iter(
        _pandoc_math_elements(
            [row for rows in formula_rows for row in rows]
        )
    )
    for formula_paragraph, rows, formula_size in zip(
        marker_paragraphs, formula_rows, formula_sizes
    ):
        row_math_elements = [next(math_elements) for _ in rows]
        number_element = formula_paragraph._p.getnext()
        number_text = ""
        if number_element is not None and number_element.tag == qn("w:p"):
            number_text = "".join(
                text.text or "" for text in number_element.iter(qn("w:t"))
            ).strip()

        # Do not leave an unnumbered Markdown display formula as literal LaTeX.
        # If no valid （章-序号） paragraph follows, render native OMML centered in
        # the existing paragraph and preserve the following content unchanged.
        if not FORMULA_NUMBER_RE.fullmatch(number_text):
            formula_paragraph.clear()
            formula_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            formula_paragraph.paragraph_format.first_line_indent = Pt(0)
            formula_paragraph.paragraph_format.space_before = Pt(0)
            formula_paragraph.paragraph_format.space_after = Pt(0)
            formula_paragraph.paragraph_format.line_spacing = 1.0
            for row_index, math_element in enumerate(row_math_elements):
                if row_index:
                    formula_paragraph.add_run().add_break()
                formula_paragraph._p.append(
                    _format_math_element(math_element, formula_size)
                )
            continue

        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _remove_table_borders(table)
        # Fixed-layout readers use tblGrid, not just each cell's tcW. Keep the
        # equation column at its declared width rather than the default third.
        _set_formula_table_widths(table)

        row = table.rows[0]
        row_pr = row._tr.get_or_add_trPr()
        row_pr.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            _set_cell_margins_zero(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        equation_cell = row.cells[1]
        for row_index, math_element in enumerate(row_math_elements):
            equation_paragraph = (
                equation_cell.paragraphs[0]
                if row_index == 0
                else equation_cell.add_paragraph()
            )
            equation_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            equation_paragraph.paragraph_format.space_before = Pt(0)
            equation_paragraph.paragraph_format.space_after = Pt(0)
            equation_paragraph.paragraph_format.line_spacing = 1.0
            equation_paragraph._p.append(
                _format_math_element(math_element, formula_size)
            )

        number_paragraph = row.cells[2].paragraphs[0]
        number_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        number_paragraph.paragraph_format.space_before = Pt(0)
        number_paragraph.paragraph_format.space_after = Pt(0)
        number_paragraph.paragraph_format.line_spacing = 1.0
        number_run = number_paragraph.add_run(number_text)
        _format_run(number_run, 10.5, east_asia="宋体", bold=False)

        formula_paragraph._p.addprevious(table._tbl)
        formula_parent = formula_paragraph._p.getparent()
        formula_parent.remove(formula_paragraph._p)
        formula_parent.remove(number_element)


def format_report(input_path, output_path):
    doc = Document(input_path)
    _remove_abstract_section_break(doc)
    _remove_section_break_before(doc, "参考文献")

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as temp_file:
        temp_path = temp_file.name
    try:
        doc.save(temp_path)
        doc = Document(temp_path)
    finally:
        os.remove(temp_path)

    title = _extract_report_title(doc)
    _configure_sections(doc, title)
    _configure_front_matter(doc)
    _configure_toc(doc)
    _configure_body(doc)
    _replace_inline_formula_markers(doc)
    _replace_formula_markers(doc)
    _repair_existing_formula_tables(doc)
    # 必须在行内/独立公式替换之后再统一段内符号，避免后插入 run 或模板
    # theme 字体让个别标点、单位、数学技术符号出现字号/字体跳变。
    _normalize_inline_symbol_formatting(doc)
    _remove_first_line_indent_after_formulas(doc)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    # Windows updates through Word; Linux writes visible caches from UNO layout
    # while retaining the original Word custom TC instructions and sections.
    refresh_result = refresh_word_fields(output_path)
    if not refresh_result.get("success"):
        raise RuntimeError("目录刷新失败: " + str(refresh_result.get("error", "unknown error")))
    return output_path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", help="待格式化的 DOCX 文件")
    parser.add_argument("output", help="格式调整后的 DOCX 文件")
    args = parser.parse_args()
    output = format_report(os.path.abspath(args.input), os.path.abspath(args.output))
    print(f"格式调整完成: {output}")


if __name__ == "__main__":
    main()
