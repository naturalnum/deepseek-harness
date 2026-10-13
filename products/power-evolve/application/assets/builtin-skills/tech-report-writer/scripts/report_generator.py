#!/usr/bin/env python3
"""科技报告生成器。

加载 output/ 目录中的 JSON 数据文件，基于模板生成 .docx 报告。
所有 JSON 文件由 LLM 在 SKILL steps 中生成。
"""

import base64
import json
import os
import re
import tempfile
from functools import lru_cache
from io import BytesIO

from bs4 import BeautifulSoup
from bs4.element import NavigableString
from PIL import Image, ImageDraw
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from config import DEFAULT_PROJECT_SOURCE, REPORT_WRITING_UNIT
from engine import generate_from_template
from format_existing_report import format_report
from report_validator import validate_report
from utils import (
    format_cover_bottom_block,
    format_cooperator_names,
    get_output_path,
    get_template_path,
    load_output_json,
    save_output_json,
)


_FIGURE_NUMBER_RE = re.compile(r"图\s*([A-Z0-9]+-\d+)")
_CATALOG_EXPLANATION_LEAD_RE = re.compile(
    r"^(?P<ws>\s*)(?P<kind>图|表)\s*[A-Z0-9]+-\d+\s*[：:]?\s*"
    r"(?P<verb>应表达|应展示|应给出|给出|展示|显示|表明|说明|反映|可见|"
    r"所示|用于表达|用于展示|用于说明|用于给出)"
)


@lru_cache(maxsize=1)
def _white_figure_placeholder_data_uri() -> str:
    """生成统一白底细边框图占位 PNG，以 data URI 形式供 Word 引擎插入。"""
    width, height = 1200, 640
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((2, 2, width - 3, height - 3), outline="black", width=2)
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    encoded = base64.b64encode(buffer.getvalue()).decode("ascii")
    image.close()
    return f"data:image/png;base64,{encoded}"


def _render_white_figure_placeholders(soup: BeautifulSoup) -> bool:
    """将文本图占位标记替换为无文字白框；图题保留在其下方。"""
    changed = False
    data_uri = _white_figure_placeholder_data_uri()
    for placeholder in soup.find_all("p"):
        placeholder_text = placeholder.get_text("", strip=True)
        match = re.match(r"^【此处插入图\s*([A-Z0-9]+-\d+)】$", placeholder_text)
        if match is None:
            continue
        number = match.group(1)
        placeholder.clear()
        placeholder["style"] = "font-family:宋体;font-size:10.5pt;text-align:center;text-indent:0"
        placeholder["data-figure-placeholder"] = number
        img = soup.new_tag(
            "img",
            src=data_uri,
            width="560",
            height="300",
            alt=f"图{number}白框占位",
        )
        placeholder.append(img)
        changed = True
    return changed


def _normalize_catalog_explanation_leads(soup: BeautifulSoup) -> bool:
    """Remove figure/table numbers from explanatory prose.

    Catalog collection is intentionally reserved for real caption paragraphs.
    Prose such as ``图1-1应表达……`` or ``表2-1给出……`` is rewritten before
    DOCX conversion so it cannot look like a caption to Word, WPS or downstream
    validators.  Only the unsafe leading phrase is changed; inline HTML after
    the prefix remains untouched.
    """
    changed = False
    for paragraph in soup.find_all("p"):
        visible = paragraph.get_text("", strip=False)
        match = _CATALOG_EXPLANATION_LEAD_RE.match(visible)
        if match is None:
            continue
        kind = match.group("kind")
        verb = match.group("verb")
        subject = "该图" if kind == "图" else "该表"
        if verb == "所示":
            replacement = f"{match.group('ws')}如{subject}所示"
        elif verb in {"显示", "表明", "反映", "可见"}:
            noun = "上述趋势" if kind == "图" else "上述结果"
            replacement = f"{match.group('ws')}{noun}{verb}"
        else:
            replacement = f"{match.group('ws')}{subject}{verb}"

        remaining = match.end()
        consumed = 0
        first_target = None
        first_offset = 0
        # Rewrite across text nodes without flattening formulas/emphasis tags.
        for node in paragraph.descendants:
            if not isinstance(node, NavigableString):
                continue
            node_text = str(node)
            node_end = consumed + len(node_text)
            if first_target is None and match.start() < node_end:
                first_target = node
                first_offset = max(0, match.start() - consumed)
            if first_target is not None and consumed < remaining:
                cut_start = max(0, match.start() - consumed)
                cut_end = min(len(node_text), remaining - consumed)
                prefix = node_text[:cut_start]
                suffix = node_text[cut_end:]
                if node is first_target:
                    node.replace_with(prefix + replacement + suffix)
                else:
                    node.replace_with(prefix + suffix)
            consumed = node_end
            if consumed >= remaining:
                break
        changed = True
    return changed


def _normalize_figure_placeholder_order(content_html: str) -> str:
    """确保图占位在图题上方，兼容旧章节数据中的反向顺序。"""
    soup = BeautifulSoup(content_html, "html.parser")
    changed = False
    for placeholder in soup.find_all("p"):
        placeholder_text = placeholder.get_text("", strip=True)
        if not placeholder_text.startswith("【此处插入图"):
            continue
        placeholder_number = _FIGURE_NUMBER_RE.search(placeholder_text)
        caption = placeholder.find_previous_sibling("p")
        if caption is None:
            continue
        caption_number = _FIGURE_NUMBER_RE.match(caption.get_text("", strip=True))
        if (
            placeholder_number is not None
            and caption_number is not None
            and placeholder_number.group(1) == caption_number.group(1)
        ):
            placeholder.extract()
            caption.insert_before(placeholder)
            changed = True
    if _normalize_catalog_explanation_leads(soup):
        changed = True
    if _render_white_figure_placeholders(soup):
        changed = True
    return str(soup) if changed else content_html


def _build_body_html(chapters: list) -> str:
    """将所有章节的 HTML 内容拼接为正文 HTML。"""
    parts = []
    for ch in chapters:
        if ch.get("content_html"):
            parts.append(_normalize_figure_placeholder_order(ch["content_html"]))
    return "\n".join(parts)


def _build_appendices_html(data: dict) -> str:
    """将可选附录数据转换为带标题和分页的 HTML。"""
    parts = []
    for appendix in data.get("appendices", []):
        appendix_id = str(appendix.get("id", "")).strip().upper()
        title = str(appendix.get("title", "")).strip()
        content = str(appendix.get("content_html", "")).strip()
        if content:
            content = _normalize_figure_placeholder_order(content)
        heading = f"附录 {appendix_id} {title}".strip()
        parts.append(
            '<h1 style="font-family:黑体;font-size:15pt;'
            'margin-top:6pt;margin-bottom:6pt;line-height:1.5;'
            f'text-indent:0;text-align:center;page-break-before:always">{heading}</h1>'
        )
        if content:
            parts.append(content)
    return "\n".join(parts)


def _build_references_html(data: dict) -> str:
    """仅在存在参考文献内容时生成参考文献标题和条目。"""
    content = str(data.get("references_html", "")).strip()
    if not content:
        return ""
    return (
        '<h1 style="font-family:黑体;font-size:15pt;'
        'margin-top:6pt;margin-bottom:6pt;line-height:1.5;'
        'text-indent:0;text-align:center;page-break-before:always">参考文献</h1>\n'
        f"{content}"
    )


def _sanitize_filename(title: str) -> str:
    """从标题中提取安全的文件名片段。"""
    # 移除不允许的字符
    chars = []
    for c in title.strip():
        if c.isalnum() or c in "_-":
            chars.append(c)
        elif c in "()（）":
            chars.append(c)
        else:
            chars.append("_")
    name = "".join(chars)
    # 清理多余的下划线
    while "__" in name:
        name = name.replace("__", "_")
    return name.strip("_")[:50] or "科技报告"


def generate_report() -> dict:
    """
    生成科技报告 Word 文档。

    从 output/ 目录读取所有 JSON 数据文件，
    构建 content_map 并调用引擎生成报告。

    Returns:
        {"success": bool, "output": str, "errors": [...]}
    """
    output_dir = os.path.dirname(get_output_path(""))
    os.makedirs(output_dir, exist_ok=True)

    # ============================================================
    # 加载所有 JSON 数据文件
    # ============================================================
    required_files = {
        "cover_info.json": "封面信息",
    }
    optional_files = {
        "outline.json": "大纲",
        "abstract.json": "摘要",
        "appendices.json": "附录",
        "references.json": "参考文献",
    }
    # 动态查找章节文件
    chapter_files = {}
    if os.path.exists(output_dir):
        for fname in sorted(os.listdir(output_dir)):
            if fname.startswith("chapter_") and fname.endswith(".json"):
                chapter_files[fname] = f"章节内容"

    loaded = {}
    missing = []
    json_errors = []

    for fname, desc in required_files.items():
        try:
            loaded[fname] = load_output_json(fname)
        except FileNotFoundError:
            missing.append(f"{fname}（{desc}）")
        except json.JSONDecodeError as e:
            json_errors.append(f"{fname} JSON 格式错误: {e}")

    for fname, desc in optional_files.items():
        try:
            loaded[fname] = load_output_json(fname)
        except FileNotFoundError:
            pass  # 可选文件
        except json.JSONDecodeError as e:
            json_errors.append(f"{fname} JSON 格式错误: {e}")

    # 加载章节文件
    chapters = []
    for fname in sorted(chapter_files.keys()):
        try:
            chapters.append(load_output_json(fname))
        except FileNotFoundError:
            pass
        except json.JSONDecodeError as e:
            json_errors.append(f"{fname} JSON 格式错误: {e}")

    if missing:
        return {
            "success": False,
            "output": None,
            "errors": [f"缺少必需文件: {', '.join(missing)}"],
        }

    if json_errors:
        return {
            "success": False,
            "output": None,
            "errors": json_errors,
        }

    # ============================================================
    # 提取数据
    # ============================================================
    cover = loaded.get("cover_info.json", {})
    missing_cover = [key for key in ("title_cn", "title_en", "project_name", "report_date")
                     if not isinstance(cover.get(key), str) or not cover[key].strip()]
    if "project_source" not in cover or not isinstance(cover["project_source"], str):
        missing_cover.append("project_source")
    if missing_cover:
        return {"success": False, "output": None,
                "errors": ["封面缺少必需数据: " + ", ".join(missing_cover)]}
    abstract_data = loaded.get("abstract.json", {})
    ref_data = loaded.get("references.json", {})
    appendices_data = loaded.get("appendices.json", {})
    outline = loaded.get("outline.json", {})

    # ============================================================
    # 构建正文 HTML
    # ============================================================
    body_html = _build_body_html(chapters)
    appendices_html = _build_appendices_html(appendices_data)
    references_html = _build_references_html(ref_data)
    back_matter_html = "\n".join(
        part for part in (appendices_html, references_html) if part
    )

    # ============================================================
    # 构建 content_map
    # ============================================================
    content_map: dict = {
        # 封面
        "REPORT_NO": cover.get("report_no", ""),
        "CLASS_NO": cover.get("class_no", ""),
        "SECRET_LEVEL": cover.get("secret_level", ""),
        "PROJECT_SOURCE": cover.get(
            "project_source", DEFAULT_PROJECT_SOURCE
        ),
        "TITLE_CN": cover.get("title_cn", ""),
        "TITLE_EN": cover.get("title_en", ""),
        "PROJECT_NAME": cover.get("project_name", ""),
        "SERIES_NO": cover.get("series_no", ""),
        "COVER_BOTTOM_BLOCK": format_cover_bottom_block(
            cover.get("cooperator", ""),
            cover.get("report_date", ""),
            REPORT_WRITING_UNIT,
        ),
        "COOPERATOR": format_cooperator_names(cover.get("cooperator", "")),
        "REPORT_DATE": cover.get("report_date", ""),
        "WORK_PERIOD": cover.get("work_period", ""),
        "LEADER": cover.get("leader", ""),
        "STAFF": cover.get("staff", ""),
        "WRITER": cover.get("writer", ""),
        # 摘要
        "ABSTRACT_CN": abstract_data.get("abstract_cn", ""),
        "KEYWORDS_CN": abstract_data.get("keywords_cn", ""),
        "ABSTRACT_EN": abstract_data.get("abstract_en", ""),
        "KEYWORDS_EN": abstract_data.get("keywords_en", ""),
        # 正文（HTML 格式）
        "BODY": body_html,
        # 附录和参考文献（均为可选）
        "APPENDICES": appendices_html,
        "REFERENCES": references_html,
        "BACK_MATTER": back_matter_html,
    }

    # ============================================================
    # 生成报告
    # ============================================================
    template_path = get_template_path("report.docx")

    # 生成文件名
    title_slug = _sanitize_filename(cover.get("title_cn", "科技报告"))
    output_filename = f"{title_slug}_科技报告.docx"
    output_path = os.path.join(output_dir, output_filename)

    result = generate_from_template(
        template_path=template_path,
        content_map=content_map,
        output_path=output_path,
    )
    if any(warning.startswith("[图片]") for warning in result.get("warnings", [])):
        return {"success": False, "output": None, "errors": ["提供的图片未能嵌入 Word，生成已阻断"]}

    if result.get("success") and result.get("output"):
        generated_path = os.path.abspath(str(result["output"]))
        # Preserve a semantic source anchor even when the legal source is blank
        # or does not contain a particular project-category phrase.
        generated_doc = Document(generated_path)
        title_index = next((i for i, p in enumerate(generated_doc.paragraphs)
                            if p.text.strip() == cover["title_cn"].strip()), None)
        if title_index is None or title_index == 0:
            return {"success": False, "output": generated_path,
                    "errors": ["模板封面缺少中文报告标题位置"]}
        source = generated_doc.paragraphs[title_index - 1]
        if source.text.strip() != cover["project_source"].strip():
            return {"success": False, "output": generated_path,
                    "errors": ["模板封面项目来源位置与输入不一致"]}
        bookmark_id = str(max([int(b.get(qn("w:id"), "0"))
            for b in generated_doc.element.iter(qn("w:bookmarkStart"))] + [0]) + 1)
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), bookmark_id)
        start.set(qn("w:name"), "PEReportSource")
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), bookmark_id)
        source._p.append(start)
        source._p.append(end)
        generated_doc.save(generated_path)
        with tempfile.NamedTemporaryFile(
            suffix=".docx",
            dir=os.path.dirname(generated_path),
            delete=False,
        ) as temp_file:
            normalized_path = temp_file.name
        try:
            format_report(generated_path, normalized_path)
            os.replace(normalized_path, generated_path)
        except Exception as exc:
            return {
                "success": False,
                "output": generated_path,
                "errors": [f"最终排版归一化失败: {exc}"],
            }
        finally:
            if os.path.exists(normalized_path):
                os.remove(normalized_path)

        validation = validate_report(generated_path)
        if not validation.get("valid"):
            return {
                "success": False,
                "output": result.get("output"),
                "errors": [
                    f"最终格式校验失败: {error}"
                    for error in validation.get("errors", [])
                ],
            }

    return {
        "success": result.get("success", False),
        "output": result.get("output"),
        "errors": result.get("errors", []),
    }
