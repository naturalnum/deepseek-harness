"""使用封面正文层中的无边框浮动文本框排版单位和发布日期。"""

from lxml import etree

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from config import (
    COVER_BOTTOM_BLOCK_WIDTH_MM,
    COVER_BOTTOM_FONT_SIZE_PT,
    COVER_HEITI_FONT,
    REPORT_WRITING_UNIT,
)

VML_NS = "urn:schemas-microsoft-com:vml"
OFFICE_NS = "urn:schemas-microsoft-com:office:office"
W10_NS = "urn:schemas-microsoft-com:office:word"
COVER_BOTTOM_TEXTBOX_ID = "CoverBottomTextBox"


def _set_attribute(element, name, value):
    element.set(qn(f"w:{name}"), str(value))


def _get_or_add(parent, tag):
    element = parent.find(qn(f"w:{tag}"))
    if element is None:
        element = OxmlElement(f"w:{tag}")
        parent.append(element)
    return element


def _remove_element(element):
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _set_paragraph_text(paragraph, text):
    paragraph.clear()
    paragraph.add_run(text)


def _format_bottom_paragraph(paragraph, keep_with_next=False):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_before = Pt(0)
    paragraph_format.space_after = Pt(0)
    paragraph_format.line_spacing = 1.0
    paragraph_format.keep_together = True
    paragraph_format.keep_with_next = keep_with_next

    p_pr = paragraph._p.get_or_add_pPr()
    frame = p_pr.find(qn("w:framePr"))
    if frame is not None:
        p_pr.remove(frame)
    snap_to_grid = _get_or_add(p_pr, "snapToGrid")
    _set_attribute(snap_to_grid, "val", "0")

    for run in paragraph.runs:
        if not run.text:
            continue
        run.font.size = Pt(COVER_BOTTOM_FONT_SIZE_PT)
        run.font.bold = False
        r_pr = run._element.get_or_add_rPr()
        fonts = _get_or_add(r_pr, "rFonts")
        for name in ("eastAsia", "ascii", "hAnsi", "cs"):
            _set_attribute(fonts, name, COVER_HEITI_FONT)


def _format_anchor_paragraph(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_before = Pt(0)
    paragraph_format.space_after = Pt(0)
    paragraph_format.line_spacing = Pt(1)

    p_pr = paragraph._p.get_or_add_pPr()
    spacing = _get_or_add(p_pr, "spacing")
    _set_attribute(spacing, "before", "0")
    _set_attribute(spacing, "after", "0")
    _set_attribute(spacing, "line", "20")
    _set_attribute(spacing, "lineRule", "exact")
    snap_to_grid = _get_or_add(p_pr, "snapToGrid")
    _set_attribute(snap_to_grid, "val", "0")


def _split_text_lines(paragraphs):
    lines = []
    for paragraph in paragraphs:
        lines.extend(
            line.strip() for line in paragraph.text.splitlines() if line.strip()
        )
    return lines


def _find_floating_bottom_table(doc):
    for table in doc.tables:
        positioning = table._tbl.tblPr.find(qn("w:tblpPr"))
        if (
            positioning is not None
            and positioning.get(qn("w:tblpYSpec")) == "bottom"
        ):
            return table
    return None


def _find_cover_bottom_textbox(doc):
    for paragraph_element in doc.element.body.findall(qn("w:p")):
        for shape in paragraph_element.iter(f"{{{VML_NS}}}shape"):
            if shape.get("id") == COVER_BOTTOM_TEXTBOX_ID:
                return paragraph_element, shape
    return None, None


def _textbox_paragraphs(doc, shape):
    content = shape.find(f".//{qn('w:txbxContent')}")
    if content is None:
        return []
    return [
        Paragraph(paragraph_element, doc._body)
        for paragraph_element in content.findall(qn("w:p"))
    ]


def _cover_section_end_element(doc):
    paragraph = next(
        (
            paragraph
            for paragraph in doc.paragraphs
            if (
                paragraph._p.pPr is not None
                and paragraph._p.pPr.find(qn("w:sectPr")) is not None
            )
        ),
        None,
    )
    return paragraph._element if paragraph is not None else doc.element.body.sectPr


def _remove_trailing_empty_cover_paragraphs(section_end):
    current = section_end.getprevious()
    while current is not None and current.tag == qn("w:p"):
        text = "".join(
            element.text or "" for element in current.iter(qn("w:t"))
        ).strip()
        has_content = (
            bool(text)
            or current.find(f".//{qn('w:drawing')}") is not None
            or current.find(f".//{qn('w:pict')}") is not None
            or current.find(f".//{qn('w:object')}") is not None
        )
        if has_content:
            break
        previous = current.getprevious()
        _remove_element(current)
        current = previous


def _append_textbox_paragraphs(doc, content, lines):
    paragraphs = []
    for line in lines:
        paragraph_element = OxmlElement("w:p")
        content.append(paragraph_element)
        paragraph = Paragraph(paragraph_element, doc._body)
        _set_paragraph_text(paragraph, line)
        paragraphs.append(paragraph)

    for index, paragraph in enumerate(paragraphs):
        _format_bottom_paragraph(
            paragraph,
            keep_with_next=index < len(paragraphs) - 1,
        )
    return paragraphs


def _create_cover_bottom_textbox(doc, lines):
    section_end = _cover_section_end_element(doc)
    if section_end is None:
        raise ValueError("文档中未找到封面分节符")
    _remove_trailing_empty_cover_paragraphs(section_end)

    anchor_element = OxmlElement("w:p")
    section_end.addprevious(anchor_element)
    anchor = Paragraph(anchor_element, doc._body)
    _format_anchor_paragraph(anchor)

    run = OxmlElement("w:r")
    anchor_element.append(run)
    pict = OxmlElement("w:pict")
    run.append(pict)

    namespace_map = {"v": VML_NS, "o": OFFICE_NS, "w10": W10_NS}
    shape_type = etree.SubElement(
        pict,
        f"{{{VML_NS}}}shapetype",
        nsmap=namespace_map,
    )
    shape_type.set("id", "_x0000_t202")
    shape_type.set("coordsize", "21600,21600")
    shape_type.set(f"{{{OFFICE_NS}}}spt", "202")
    shape_type.set("path", "m,l,21600r21600,l21600,xe")
    stroke = etree.SubElement(shape_type, f"{{{VML_NS}}}stroke")
    stroke.set("joinstyle", "miter")
    path = etree.SubElement(shape_type, f"{{{VML_NS}}}path")
    path.set("gradientshapeok", "t")
    path.set(f"{{{OFFICE_NS}}}connecttype", "rect")

    shape = etree.SubElement(
        pict,
        f"{{{VML_NS}}}shape",
        nsmap=namespace_map,
    )
    shape.set("id", COVER_BOTTOM_TEXTBOX_ID)
    shape.set(f"{{{OFFICE_NS}}}spid", "_x0000_s2049")
    shape.set("type", "#_x0000_t202")
    width_pt = COVER_BOTTOM_BLOCK_WIDTH_MM * 72 / 25.4
    shape.set(
        "style",
        (
            "position:absolute;"
            f"width:{width_pt:.3f}pt;"
            "height:1pt;"
            "z-index:251659264;"
            "mso-position-horizontal:center;"
            "mso-position-horizontal-relative:margin;"
            "mso-position-vertical:bottom;"
            "mso-position-vertical-relative:margin;"
            "mso-fit-shape-to-text:t;"
            "mso-wrap-style:none"
        ),
    )
    shape.set("filled", "f")
    shape.set("stroked", "f")

    textbox = etree.SubElement(shape, f"{{{VML_NS}}}textbox")
    textbox.set("inset", "0,0,0,0")
    textbox.set("style", "mso-fit-shape-to-text:t")
    content = OxmlElement("w:txbxContent")
    textbox.append(content)
    paragraphs = _append_textbox_paragraphs(doc, content, lines)

    wrap = etree.SubElement(shape, f"{{{W10_NS}}}wrap")
    wrap.set("type", "none")
    return anchor, paragraphs


def normalize_cover_bottom_template(doc):
    """将旧页底结构迁移为封面正文层中的无边框浮动文本框。"""
    anchor_element, shape = _find_cover_bottom_textbox(doc)
    if shape is not None:
        lines = _split_text_lines(_textbox_paragraphs(doc, shape))
        _remove_element(anchor_element)
        _, paragraphs = _create_cover_bottom_textbox(
            doc,
            lines or ["{{COVER_BOTTOM_BLOCK}}"],
        )
        return paragraphs

    lines = extract_and_remove_cover_bottom(doc)
    _, paragraphs = _create_cover_bottom_textbox(
        doc,
        lines or ["{{COVER_BOTTOM_BLOCK}}"],
    )
    return paragraphs


def extract_and_remove_cover_bottom(doc):
    """读取并移除页底结构，兼容文本框、表格、页脚和普通段落。"""
    anchor_element, shape = _find_cover_bottom_textbox(doc)
    if shape is not None:
        lines = _split_text_lines(_textbox_paragraphs(doc, shape))
        _remove_element(anchor_element)
        return lines

    table = _find_floating_bottom_table(doc)
    if table is not None:
        lines = _split_text_lines(table.cell(0, 0).paragraphs)
        _remove_element(table._tbl)
        return lines

    footer_lines = _split_text_lines(doc.sections[0].footer.paragraphs)
    if footer_lines and footer_lines[0] == REPORT_WRITING_UNIT:
        for paragraph in doc.sections[0].footer.paragraphs:
            paragraph.clear()
        return footer_lines

    framed = [
        paragraph
        for paragraph in doc.paragraphs
        if (
            paragraph._p.pPr is not None
            and paragraph._p.pPr.find(qn("w:framePr")) is not None
        )
    ]
    if framed:
        lines = _split_text_lines(framed)
        for paragraph in framed:
            _remove_element(paragraph._element)
        return lines

    anchor_index = next(
        (
            index
            for index, paragraph in enumerate(doc.paragraphs)
            if (
                paragraph.text.strip() == "{{COVER_BOTTOM_BLOCK}}"
                or paragraph.text.strip().startswith(REPORT_WRITING_UNIT)
            )
        ),
        None,
    )
    if anchor_index is None:
        raise ValueError("文档中未找到封面页底区块")

    candidates = []
    for paragraph in doc.paragraphs[anchor_index:]:
        p_pr = paragraph._p.pPr
        if p_pr is not None and p_pr.find(qn("w:sectPr")) is not None:
            break
        if paragraph.text.strip():
            candidates.append(paragraph)
    lines = _split_text_lines(candidates)
    for paragraph in candidates:
        _remove_element(paragraph._element)
    return lines


def set_cover_bottom_textbox(doc, lines):
    """创建无边框浮动文本框，每个单位和日期分别使用独立段落。"""
    normalized = [str(line).strip() for line in lines if str(line).strip()]
    if len(normalized) < 2:
        raise ValueError("封面页底至少应包含报告编写单位和发布日期")

    anchor_element, _ = _find_cover_bottom_textbox(doc)
    if anchor_element is not None:
        _remove_element(anchor_element)
    table = _find_floating_bottom_table(doc)
    if table is not None:
        _remove_element(table._tbl)
    _, paragraphs = _create_cover_bottom_textbox(doc, normalized)
    return paragraphs
