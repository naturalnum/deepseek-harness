"""Real CLI/apply/confirm and unchanged Word consumers with synthetic scoped source files.

Test responses and confirmations are interface fixtures, never real model/reviewer acceptance.
Run with the existing report Python test dependencies (docx, bs4, jsonschema).
"""
import copy
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[1]
SKILLS = ROOT.parent
sys.path.insert(0, str(ROOT / "scripts"))
import workflow_engine as workflow

REGISTRY = json.loads((ROOT / "skill-registry.json").read_text())
FIXTURE = ROOT / "tests/fixtures/chapter-review-passed.json"


class ReferenceEvidence(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.work = Path(self.directory.name).resolve()
        self.source = self.work / "source"; self.source.mkdir()
        self.state = self.work / "state.json"
        self.content = {"internet_access": "unavailable", "collection_confirmed": False,
                        "organization_approach": {"scope": "合成材料的证据范围"}, "items": []}
        for index, (name, marker) in enumerate((("task-book.md", "R"), ("experiment.csv", "DS"),
                                                ("method.md", "R"), ("output/verification/verification.log", "R")), 1):
            path = self.source / name; path.parent.mkdir(parents=True, exist_ok=True)
            raw = f"synthetic source {index}\n".encode(); path.write_bytes(raw)
            digest = hashlib.sha256(raw).hexdigest()
            self.content["items"].append({
                "id": f"REF-{index}", "title": f"合成材料{index}",
                "reference_type": "dataset" if marker == "DS" else "report",
                "source_kind": "local_file", "verification_status": "verified",
                "supports": ["本接口测试的材料字节"], "doi_or_identifier": "sha256:" + digest,
                "local_file_evidence": {"relative_path": name, "sha256": digest, "byte_size": len(raw)},
                "gbt7714_entry": f"测试组. 合成材料{index}[{marker}]. 2026.",
            })
        state = json.loads(FIXTURE.read_text())
        state["stage"] = "OUTLINE_CONFIRMED"; state["pending_action"] = None
        state["artifacts"]["reference_plan"] = {"versions": [], "approved_version": None}
        state["confirmations"].pop("reference_plan", None)
        self.write(self.state, state)

    def write(self, path, data):
        path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n")

    def cli(self, *args):
        return subprocess.run([sys.executable, str(ROOT / "scripts/workflow_cli.py"), *map(str, args)],
                              capture_output=True, text=True, timeout=10,
                              env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"})

    def succeed(self, *args):
        result = self.cli(*args)
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        return result

    def reject(self, *args):
        before = self.state.read_bytes()
        result = self.cli(*args)
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(self.state.read_bytes(), before)
        return result

    def request(self):
        path = self.work / "request.json"
        self.succeed("next", self.state, "--request-out", path)
        return json.loads(path.read_text())

    def response(self, request, content=None, **extra):
        path = self.work / "response.json"
        self.write(path, {"request_id": request["orchestration"]["request_id"],
                          "skill_name": request["skill_name"], "action": request["action"], "status": "success",
                          "artifact": {"artifact_type": "reference_plan", "artifact_id": "refplan-synthetic",
                                       "version": 1, "content": self.content if content is None else content, **extra}})
        return path

    def apply_local(self):
        request = self.request()
        self.assertEqual(request["payload"]["reference_evidence_contract"], 2)
        self.succeed("apply", self.state, self.response(request), "--source-root", self.source)
        return json.loads(self.state.read_text())["artifacts"]["reference_plan"]["versions"][-1]

    def test_real_apply_and_confirm_require_explicit_root_and_store_checked_contract(self):
        request = self.request(); path = self.response(request)
        self.reject("apply", self.state, path)
        self.succeed("apply", self.state, path, "--source-root", self.source)
        state = json.loads(self.state.read_text())
        self.assertEqual(state["stage"], "REFERENCE_PLAN_DRAFTED")
        self.assertIsNone(state["pending_action"])
        self.assertEqual(state["artifacts"]["reference_plan"]["versions"][-1]["reference_evidence_contract"], 2)
        self.assertNotIn(str(self.source), self.state.read_text())
        self.reject("confirm", self.state, "reference_plan")
        self.succeed("confirm", self.state, "reference_plan", "--source-root", self.source)
        state = json.loads(self.state.read_text())
        self.assertEqual(state["stage"], "REFERENCE_PLAN_CONFIRMED")
        self.assertEqual(state["artifacts"]["reference_plan"]["approved_version"], 1)

    def test_model_cannot_omit_new_contract_structure_or_return_aliases(self):
        request = self.request()
        legacy = json.loads(FIXTURE.read_text())["artifacts"]["reference_plan"]["versions"][-1]["content"]
        for content in (legacy, {"organization_approach": {}}, {"reference_registry": self.content},
                        {**self.content, "references": []}):
            with self.subTest(keys=list(content)):
                self.reject("apply", self.state, self.response(request, content), "--source-root", self.source)
        for marker in (1, True, "2"):
            self.reject("apply", self.state, self.response(request, reference_evidence_contract=marker),
                        "--source-root", self.source)

    def test_changed_deleted_and_linked_sources_cannot_inherit_apply_validation_at_confirm(self):
        self.apply_local()
        path = self.source / "method.md"; raw = path.read_bytes()
        path.write_bytes(raw + b"x")
        self.reject("confirm", self.state, "reference_plan", "--source-root", self.source)
        path.unlink()
        self.reject("confirm", self.state, "reference_plan", "--source-root", self.source)
        outside = self.work / "outside.md"; outside.write_bytes(raw); path.symlink_to(outside)
        self.reject("confirm", self.state, "reference_plan", "--source-root", self.source)
        path.unlink(); path.write_bytes(raw)
        self.succeed("confirm", self.state, "reference_plan", "--source-root", self.source)

    def test_reserved_local_metadata_outside_items_cannot_bypass_apply_or_confirm(self):
        request = self.request(); pending = json.loads(self.state.read_text())
        for location in ("organization", "artifact", "organization-root", "artifact-root", "placeholder"):
            content = {"internet_access": "unavailable", "items": []}
            extra = {}
            metadata = copy.deepcopy(self.content["items"][0])
            if location == "organization":
                content["organization_approach"] = {"sources": [metadata]}
            elif location == "artifact":
                extra["local_file_evidence"] = metadata["local_file_evidence"]
            elif location == "organization-root":
                content["organization_approach"] = {"source_root": str(self.source)}
            elif location == "artifact-root":
                extra["source_root"] = str(self.source)
            else:
                content["organization_approach"] = {"sources": [{
                    "official_url": "https://local.invalid/material",
                    "verification_sources": ["https://local.invalid/material"]}]}
            response = self.response(request, content, **extra)
            self.reject("apply", self.state, response)
            draft = copy.deepcopy(pending); draft.update(stage="REFERENCE_PLAN_DRAFTED", pending_action=None)
            artifact = json.loads(response.read_text())["artifact"]
            artifact["reference_evidence_contract"] = 2
            draft["artifacts"]["reference_plan"]["versions"] = [artifact]
            self.write(self.state, draft); self.reject("confirm", self.state, "reference_plan")
            self.write(self.state, pending)
        # Mentioning a local-file identifier in prose does not declare an unvalidated source.
        content = {"internet_access": "unavailable", "items": [],
                   "organization_approach": {"description": "local-file: is not used as an external DOI"}}
        self.succeed("apply", self.state, self.response(request, content))
        self.succeed("confirm", self.state, "reference_plan")

    def test_old_requests_and_drafts_cannot_hide_local_fields_in_legacy_aliases(self):
        request = self.request()
        state = json.loads(self.state.read_text()); state["pending_action"]["payload"].pop("reference_evidence_contract")
        self.write(self.state, state)
        legacy = copy.deepcopy(json.loads(FIXTURE.read_text())["artifacts"]["reference_plan"]["versions"][-1]["content"])
        for alias in ("references", "verified_references"):
            for extra in ({"source_kind": "local"}, {"local_file_evidence": self.content["items"][0]["local_file_evidence"]},
                          {"relative_path": "method.md", "sha256": "0" * 64, "byte_size": 1},
                          {"doi_or_identifier": "local-file:test-material.md"}):
                with self.subTest(alias=alias, extra=list(extra)):
                    content = copy.deepcopy(legacy); content.setdefault(alias, [{}])[0].update(extra)
                    self.reject("apply", self.state, self.response(request, content), "--source-root", self.source)
                    draft = copy.deepcopy(state); draft["pending_action"] = None; draft["stage"] = "REFERENCE_PLAN_DRAFTED"
                    draft["artifacts"]["reference_plan"]["versions"] = [
                        {"artifact_type": "reference_plan", "artifact_id": "old", "version": 1, "content": content}]
                    self.write(self.state, draft)
                    self.reject("confirm", self.state, "reference_plan", "--source-root", self.source)
                    self.write(self.state, state)

    def test_captured_old_nonlocal_states_and_pending_replay_keep_the_existing_shape(self):
        for name in ("chapter-review-passed.json", "chapter-revision-limit.json"):
            state = json.loads((ROOT / "tests/fixtures" / name).read_text())
            state.update(stage="REFERENCE_PLAN_DRAFTED", pending_action=None)
            state["confirmations"].pop("reference_plan", None)
            state["artifacts"]["reference_plan"]["approved_version"] = None
            self.write(self.state, state); self.succeed("confirm", self.state, "reference_plan")
            state.update(stage="OUTLINE_CONFIRMED", pending_action=None)
            old = copy.deepcopy(state["artifacts"]["reference_plan"]["versions"][-1])
            state["artifacts"]["reference_plan"] = {"versions": [], "approved_version": None}
            state["confirmations"].pop("reference_plan", None); self.write(self.state, state)
            request = self.request()
            state = json.loads(self.state.read_text()); state["pending_action"]["payload"].pop("reference_evidence_contract")
            self.write(self.state, state)
            self.succeed("apply", self.state, self.response(request, old["content"]))
            result = json.loads(self.state.read_text())["artifacts"]["reference_plan"]["versions"][-1]
            self.assertNotIn("reference_evidence_contract", result)
            self.assertEqual(result["content"], old["content"])

    def test_legacy_placeholder_http_cannot_be_applied_or_confirmed(self):
        request = self.request()
        state = json.loads(self.state.read_text()); state["pending_action"]["payload"].pop("reference_evidence_contract")
        self.write(self.state, state)
        legacy = json.loads(FIXTURE.read_text())["artifacts"]["reference_plan"]["versions"][-1]["content"]
        for url in ("https://local.invalid/source", "https://EXAMPLE.INVALID./source"):
            content = copy.deepcopy(legacy)
            content["references"][0].update(official_url="https://www.itl.nist.gov/",
                                             verification_sources=[url])
            content["references"][0]["gbt7714"] += " " + url + "."
            self.reject("apply", self.state, self.response(request, content))
            draft = copy.deepcopy(state); draft.update(stage="REFERENCE_PLAN_DRAFTED", pending_action=None)
            draft["artifacts"]["reference_plan"]["versions"] = [
                {"artifact_type": "reference_plan", "artifact_id": "old", "version": 1, "content": content}]
            self.write(self.state, draft); self.reject("confirm", self.state, "reference_plan")
            self.write(self.state, state)

    def test_quality_gate_wrapper_forwards_source_root(self):
        path = self.work / "references.json"; self.write(path, self.content)
        def run(*extra):
            return subprocess.run([sys.executable, str(ROOT / "scripts/run_quality_gate.py"),
                                   "--references", str(path), *map(str, extra)], capture_output=True, text=True,
                                  timeout=10, env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"})
        self.assertEqual(run("--source-root", self.source).returncode, 0)
        self.assertEqual(run().returncode, 1)

    def test_published_schemas_preserve_explicit_local_shape_and_fixed_request_contract(self):
        from jsonschema import Draft202012Validator
        generator = SKILLS / "power-grid-technical-report-writer"
        schema = json.loads((generator / "schemas/reference-registry.schema.json").read_text())
        check = Draft202012Validator(schema)
        self.assertEqual(list(check.iter_errors(self.content)), [])
        for mutation in ("bool-size", "extra-root", "missing-kind", "unknown-kind", "online-url"):
            content = copy.deepcopy(self.content); item = content["items"][0]
            if mutation == "bool-size":
                item["local_file_evidence"]["byte_size"] = True
            elif mutation == "extra-root":
                item["local_file_evidence"]["source_root"] = str(self.source)
            elif mutation == "missing-kind":
                item.pop("source_kind")
            elif mutation == "unknown-kind":
                item["source_kind"] = "local"
            else:
                item["official_url"] = "https://www.itl.nist.gov/"
            self.assertTrue(list(check.iter_errors(content)), mutation)
        request_check = Draft202012Validator(json.loads((generator / "schemas/orchestration-request.schema.json").read_text()))
        request = self.request()
        self.assertEqual(list(request_check.iter_errors(request)), [])
        request["payload"].pop("reference_evidence_contract")
        self.assertTrue(list(request_check.iter_errors(request)))
        for value in (1, True, "2"):
            request["payload"]["reference_evidence_contract"] = value
            self.assertTrue(list(request_check.iter_errors(request)))

    def test_chapter_and_formal_review_payloads_preserve_the_same_top_level_items(self):
        artifact = self.apply_local()
        self.succeed("confirm", self.state, "reference_plan", "--source-root", self.source)
        original = json.loads(self.state.read_text())
        for stage in ("REFERENCE_PLAN_CONFIRMED", "CHAPTER_MATERIALS_READY", "CHAPTER_PLAN_CONFIRMED",
                      "CHAPTER_DRAFTED", "CHAPTER_REVISION_REQUIRED", "REPORT_ASSEMBLED"):
            state = copy.deepcopy(original); state["stage"] = stage; state["current_chapter_id"] = "1"
            if stage == "REFERENCE_PLAN_CONFIRMED":
                state["chapters"]["1"]["materials"] = None
            request = workflow.next_action(state, REGISTRY)
            self.assertEqual(request["payload"]["reference_plan"], artifact)
            self.assertEqual(request["payload"]["reference_plan"]["content"]["items"], self.content["items"])
        self.assertEqual(workflow._assembly_payload(original)["reference_plan"], artifact)
        self.assertEqual(workflow._word_format_payload(original)["reference_plan"], artifact)

    def test_actual_unchanged_formatter_reads_items_and_final_markdown_into_docx(self):
        artifact = self.apply_local()
        sys.path.insert(0, str(SKILLS / "tech-report-writer/scripts"))
        from orchestration_adapter import references_data, final_report_data
        from main import _validate_references_data
        from engine.parser import parse_html_to_elements
        from docx import Document
        from report_validator import _gbt7714_reference_error
        entries = [item["gbt7714_entry"] for item in artifact["content"]["items"]]
        prepared = references_data(artifact)
        self.assertTrue(_validate_references_data(prepared)[0])
        self.assertEqual(prepared["references_html"].count("<p "), len(entries))
        for entry in entries:
            self.assertIn(entry, prepared["references_html"])
            self.assertIsNone(_gbt7714_reference_error(entry))
        source = ("# 合成引用接口测试报告\n\n## 中文摘要\n\n合成接口检查。\n\n**关键词：** 电网；数据\n\n"
                  "## ABSTRACT\n\nSynthetic interface check.\n\n**Keywords:** grid; data\n\n"
                  "# 1 项目概述\n\n只验证材料字节和引用传递。\n\n# 参考文献\n\n"
                  + "\n".join(f"[{i}] {entry}" for i, entry in enumerate(entries, 1)))
        final = final_report_data({"content": {"markdown": source}})
        self.assertTrue(_validate_references_data(final["references"])[0])
        doc = Document()
        _, errors, _ = parse_html_to_elements(doc, final["references"]["references_html"], "REFERENCES")
        self.assertEqual(errors, [])
        path = self.work / "local-reference-consumer.docx"; doc.save(path)
        self.assertEqual([p.text for p in Document(path).paragraphs],
                         [f"[{i}] {entry}" for i, entry in enumerate(entries, 1)])
        # The final Markdown is authoritative; do not silently restore an old registry's text.
        changed = source.replace(entries[0], entries[0].replace("合成材料1", "明确更新后的题名"))
        changed_final = final_report_data({"content": {"markdown": changed}})
        self.assertIn("明确更新后的题名", changed_final["references"]["references_html"])
        self.assertNotIn(entries[0], changed_final["references"]["references_html"])


if __name__ == "__main__":
    unittest.main()
